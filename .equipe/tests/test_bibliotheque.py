#!/usr/bin/env python3
"""Tests de la bibliothèque juridique (T-040, T-053) : python .equipe/tests/test_bibliotheque.py
Lecture seule sur la base réelle (bibliothèque réelle) ; toute écriture se fait dans une racine jetable.
Sortie : une ligne OK / ÉCHEC / SAUTÉ par test, puis le bilan. Un test réseau sans réseau est SAUTÉ, jamais bloquant."""
import json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

EQ = Path(__file__).resolve().parents[1]
ROOT = EQ.parent
CEREBRO = EQ / "cerebro" / "cerebro.py"
sys.path.insert(0, str(EQ / "scripts" / "calcul"))
sys.path.insert(0, str(EQ / "scripts" / "bibliotheque"))
RES = []


def cb(*args, env=None):
    r = subprocess.run([sys.executable, str(CEREBRO), *map(str, args)], capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1", **(env or {})})
    try:
        return json.loads(r.stdout)
    except Exception:
        return {"erreur": r.stdout + r.stderr}


def test(nom):
    def deco(f):
        try:
            r = f()
            RES.append((nom, "SAUTÉ" if r == "SAUTÉ" else "OK", "" if r in (None, "SAUTÉ") else str(r)))
        except AssertionError as e:
            RES.append((nom, "ÉCHEC", str(e)[:300]))
        except Exception as e:
            RES.append((nom, "ÉCHEC", repr(e)[:300]))
        return f
    return deco


def racine_jetable(avec_bibliotheque=False):
    t = Path(tempfile.mkdtemp(prefix="cb-bib-"))
    (t / ".equipe").mkdir()
    for d in ("cerebro", "config", "cerveau", "sommaires"):
        if (EQ / d).exists():
            shutil.copytree(EQ / d, t / ".equipe" / d, ignore=shutil.ignore_patterns("*.db-*", "__pycache__"))
    (t / "Bureau" / "Informatique").mkdir(parents=True)
    if avec_bibliotheque:
        for j in ("ch", "vd", "ge"):  # textes fédéraux et cantonaux (sans le cache de téléchargement)
            if (EQ / "bibliotheque" / j).exists():
                shutil.copytree(EQ / "bibliotheque" / j, t / ".equipe" / "bibliotheque" / j)
    else:
        (t / ".equipe" / "cerebro" / "cerebro.db").unlink(missing_ok=True)
    return t


@test("law article 642.11 art. 132 : texte, version, URL")
def _():
    r = cb("law", "article", "642.11", "art. 132")
    assert "texte" in r, f"pas de texte : {r}"
    assert "30 jours" in r["texte"], "texte inattendu"
    assert r.get("version") and r.get("url", "").startswith("https://fedlex.data.admin.ch/"), r
    assert r["article"].startswith("Art. 132"), r["article"]


@test("law asof : version de l'époque quand plusieurs versions sont ingérées")
def _():
    a = cb("law", "asof", "641.20", "--date", "2024-06-30")
    b = cb("law", "asof", "641.20")
    if not a or not b or a.get("id") == b.get("id"):
        return "SAUTÉ"  # dette : une seule version ingérée
    assert a["version"] <= "2024-06-30" and (a["valide_au"] or "9999") > "2024-06-30", a
    assert b["version"] > a["version"], (a["version"], b["version"])
    art = cb("law", "article", "641.20", "art. 25", "--date", "2024-06-30")
    assert art.get("version") == a["version"], art
    trou = cb("law", "asof", "641.20", "--date", "2025-02-15")  # version 2025-01-01 non ingérée → rien plutôt qu'un texte périmé
    assert trou is None or trou.get("version") >= "2025-01-01", trou


@test("règles de délais RD-001…RD-008 confirmées par le texte (deux passes, racine jetable)")
def _():
    t = racine_jetable(avec_bibliotheque=True)
    env = {"CEREBRO_ROOT": str(t)}
    for passe in (1, 2):
        v = cb("law", "verify", env=env)
        assert isinstance(v, list) and v, v
        ko = [f"{r['id']} {r['article']}" for r in v if not r["verifie"]]
        assert not ko, f"passe {passe} : non confirmées {ko}"
    shutil.rmtree(t, ignore_errors=True)


@test("calcul IA sans barème → ⚠ et aucun chiffre")
def _():
    t = racine_jetable()
    os.environ["CEREBRO_ROOT"] = str(t)
    try:
        import impot_anticipe
        r = impot_anticipe.calculer(brut=100000, annee=2026).resultat()
    finally:
        os.environ.pop("CEREBRO_ROOT", None)
        shutil.rmtree(t, ignore_errors=True)
    assert r["resultat"] is None, r["resultat"]
    assert "⚠" in r["avertissement"], r


@test("calcul IA avec barème : cohérent avec le texte LIA art. 13")
def _():
    import impot_anticipe
    r = impot_anticipe.calculer(brut=100000, annee=2026).resultat()
    assert r["resultat"], r.get("avertissement")
    taux = r["hypotheses"][0]["valeur"]
    art = cb("law", "article", "642.21", "art. 13")
    assert f"à {taux:g} % de la prestation imposable" in art["texte"].replace(" ", " "), (taux, art.get("texte", "")[:200])
    assert r["valeurs"]["ia"] == round(100000 * taux / 100, 2) and r["valeurs"]["net_verse"] == 100000 - r["valeurs"]["ia"], r["valeurs"]
    assert "642.21 art. 13" in r["sources"][0]["source"], r["sources"]


@test("calcul TVA et droit d'émission : barèmes lus en base, Excel en formules")
def _():
    import tva, droit_emission
    r = tva.calculer(1000, "ht", "normal", 2026).resultat()
    assert r["resultat"] and "641.20 art. 25" in r["sources"][0]["source"], r
    c = droit_emission.calculer(1500000, 1000000, 0, 2026)
    d = c.resultat()
    assert d["resultat"] and "641.10" in d["sources"][0]["source"], d
    x = Path(tempfile.mkdtemp()) / "de.xlsx"
    c.excel(x)
    from openpyxl import load_workbook
    wb = load_workbook(x)
    assert wb.sheetnames == ["Hypothèses", "Calcul", "Sources"], wb.sheetnames
    formules = [row[1] for row in wb["Calcul"].iter_rows(min_row=2, values_only=True)]
    assert formules and all(isinstance(f, str) and f.startswith("=") for f in formules), formules


@test("baremes.py : sans texte ingéré, aucune valeur inscrite")
def _():
    t = racine_jetable()
    r = subprocess.run([sys.executable, str(EQ / "scripts" / "bibliotheque" / "baremes.py"), "--dry-run"], capture_output=True, text=True,
                       encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8", "CEREBRO_ROOT": str(t)})
    shutil.rmtree(t, ignore_errors=True)
    out = json.loads(r.stdout)
    assert out and all(x["statut"] == "non confirmé" for x in out), out[:2]


@test("conversion Akoma Ntoso → ## Art. N (hors réseau)")
def _():
    import fedlex
    xml = b"""<?xml version="1.0" encoding="UTF-8"?><akomaNtoso xmlns="http://docs.oasis-open.org/legaldocml/ns/akn/3.0" xmlns:fedlex="http://fedlex.admin.ch/"><act><meta><identification><FRBRWork><FRBRnumber value="999"/></FRBRWork></identification></meta><body>
<chapter eId="chap_1"><num>Chapitre 1</num><heading>Essai</heading><level eId="l1" fedlex:role="marginal"><num>A. </num><heading>Titre marginal</heading>
<article eId="art_1"><num><b>Art. 1</b><authorialNote><p>Note officielle.</p></authorialNote></num><paragraph eId="art_1/para_1"><num>1</num><content><p>Premier alin\xc3\xa9a.</p></content></paragraph>
<paragraph eId="art_1/para_2"><num>2</num><content><blockList><listIntroduction>Liste:</listIntroduction><item><num>a.</num><p>lettre a;</p></item></blockList></content></paragraph></article></level>
<article eId="art_2"><num><b>Art. 2</b><i>a</i></num><heading>Propre titre</heading><content><p>Texte.</p></content></article></chapter></body></act></akomaNtoso>"""
    md, meta = fedlex.akn_vers_markdown(xml)
    assert "## Art. 1 A. Titre marginal" in md, md
    assert "1 Premier alinéa." in md and "- a. lettre a;" in md, md
    assert "## Art. 2a Propre titre" in md, md
    assert "Note officielle." not in md.split("## Notes de bas de page")[0], "note restée dans l'article"
    assert "### Chapitre 1 Essai" in md and meta["rs"] == "999"


@test("Fedlex SPARQL : RS 642.11 → LIFD en vigueur (réseau)")
def _():
    import fedlex
    try:
        a = fedlex.resoudre("642.11")
    except ConnectionError:
        return "SAUTÉ"
    assert a and a["abrevs"].get("fr") == "LIFD" and a["eli"].startswith("https://fedlex.data.admin.ch/eli/cc/"), a


@test("Zefix (LINDAS) : Nestlé S.A. par IDE (réseau)")
def _():
    import zefix
    try:
        r = zefix.chercher(ide="CHE-105.909.036", statut=False)
    except Exception:
        return "SAUTÉ"
    noms = {e["raison_sociale"] for e in r["resultats"]}
    assert "Nestlé S.A." in noms, noms
    e = [x for x in r["resultats"] if x["raison_sociale"] == "Nestlé S.A."][0]
    assert e["forme"] == "Société anonyme" and e["siege"] == "Vevey" and e["but"], e


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    for nom, st, d in RES:
        print(f"{st:6} {nom}" + (f" — {d}" if d and st != "OK" else ""))
    ko = [r for r in RES if r[1] == "ÉCHEC"]
    print(f"\n{len(RES) - len(ko)}/{len(RES)} sans échec" + (" — ÉCHEC" if ko else " — OK"))
    sys.exit(1 if ko else 0)

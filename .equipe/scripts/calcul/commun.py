"""Socle commun des calculs déterministes (§4.7) : barèmes lus dans la base via `cerebro rates get` (jamais de valeur en dur),
une seule expression par ligne de calcul, évaluée en Python pour la sortie JSON et recopiée telle quelle en formule Excel.
Barème absent → résultat ⚠ sans aucun chiffre calculé. Stdlib + openpyxl (pour --excel)."""
import datetime as dt, json, os, re, subprocess, sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

CEREBRO = Path(__file__).resolve().parents[2] / "cerebro" / "cerebro.py"
AVERT = "⚠"


def cerebro(*args):
    r = subprocess.run([sys.executable, str(CEREBRO), *[str(a) for a in args]], capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1"})
    try:
        return json.loads(r.stdout or "{}")
    except Exception:
        return {"erreur": (r.stdout or r.stderr).strip()[:300]}


def bareme(nom, cle, annee):
    """valeur la plus récemment vérifiée pour (nom, année, clé), ou None"""
    rows = cerebro("rates", "get", nom, "--annee", annee, "--cle", cle)
    if not isinstance(rows, list) or not rows:
        return None
    r = sorted(rows, key=lambda x: (x.get("verifie_le") or "", x.get("id") or ""))[-1]
    try:
        v = float(r["valeur"])
    except (TypeError, ValueError):
        return None
    return {"id": r["id"], "nom": nom, "cle": cle, "annee": int(annee), "valeur": v, "source": r.get("source") or "", "verifie_le": r.get("verifie_le")}


def ROUND(x, n=2):
    return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))


FONCTIONS = {"ROUND": ROUND, "MAX": max, "MIN": min}


class Calcul:
    """hypothèses (entrées et barèmes) + lignes de calcul (expressions à {clé}) + sources + réserves"""

    def __init__(self, titre, annee):
        self.titre, self.annee = titre, int(annee)
        self.hyp, self.lignes, self.sources, self.reserves, self.manquants = [], [], [], [], []

    def entree(self, cle, libelle, valeur, unite="CHF", origine="déclaré"):
        self.hyp.append({"cle": cle, "libelle": libelle, "valeur": valeur, "unite": unite, "origine": origine})

    def taux(self, cle, libelle, nom, cle_bareme, unite="%"):
        b = bareme(nom, cle_bareme, self.annee)
        if not b:
            self.manquants.append(f"{nom}.{cle_bareme} ({self.annee})")
            self.hyp.append({"cle": cle, "libelle": libelle, "valeur": None, "unite": unite, "origine": f"{AVERT} barème absent : {nom}.{cle_bareme} {self.annee}"})
            return None
        self.hyp.append({"cle": cle, "libelle": libelle, "valeur": b["valeur"], "unite": unite, "origine": f"barème {b['id']}"})
        self.sources.append(b)
        return b["valeur"]

    def ligne(self, cle, libelle, expr, unite="CHF"):
        self.lignes.append({"cle": cle, "libelle": libelle, "expr": expr, "unite": unite})

    def reserve(self, txt):
        self.reserves.append(txt)

    def evaluer(self):
        if self.manquants:
            return None
        env = {h["cle"]: h["valeur"] for h in self.hyp}
        res = {}
        for l in self.lignes:
            py = re.sub(r"\{(\w+)\}", lambda m: f"_v['{m.group(1)}']", l["expr"])
            v = eval(py, {"__builtins__": {}, **FONCTIONS}, {"_v": {**env, **res}})
            res[l["cle"]] = ROUND(v, 2) if l["unite"] == "CHF" else v
        return res

    def resultat(self):
        res = self.evaluer()
        out = {"calcul": self.titre, "annee": self.annee, "etabli_le": dt.date.today().isoformat(),
               "hypotheses": [{k: h[k] for k in ("libelle", "valeur", "unite", "origine")} for h in self.hyp],
               "sources": [{"bareme": s["id"], "valeur": s["valeur"], "source": s["source"], "verifie_le": s["verifie_le"]} for s in self.sources],
               "reserves": list(self.reserves)}
        if res is None:
            out["resultat"] = None
            out["avertissement"] = f"{AVERT} barème absent de la base : {', '.join(self.manquants)}. Aucun chiffre n'est calculé de mémoire ; ingérer le texte officiel puis `python baremes.py`."
        else:
            out["resultat"] = {l["libelle"]: res[l["cle"]] for l in self.lignes}
            out["valeurs"] = res
        return out

    # ------------------------------------------------------------ Excel (formules, jamais de valeurs calculées en dur)
    def excel(self, chemin):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        h = wb.active
        h.title = "Hypothèses"
        gras = Font(bold=True)
        h.append([self.titre, "", "", f"année {self.annee}", f"établi le {dt.date.today().isoformat()}"])
        h["A1"].font = Font(bold=True, size=13)
        h.append(["Hypothèse", "Valeur", "Unité", "Origine"])
        for c in h[2]:
            c.font = gras
        ref = {}
        for h_ in self.hyp:
            h.append([h_["libelle"], h_["valeur"], h_["unite"], h_["origine"]])
            ref[h_["cle"]] = f"'Hypothèses'!$B${h.max_row}"
            if h_["valeur"] is None:
                h.cell(h.max_row, 2).fill = PatternFill("solid", fgColor="FFF2CC")
        c = wb.create_sheet("Calcul")
        c.append(["Poste", "Montant", "Unité", "Formule (lecture)"])
        for x in c[1]:
            x.font = gras
        if self.manquants:
            c.append([f"{AVERT} barème absent : {', '.join(self.manquants)} — aucun calcul", None, "", ""])
        else:
            for l in self.lignes:
                f = re.sub(r"\{(\w+)\}", lambda m: ref[m.group(1)], l["expr"])
                c.append([l["libelle"], "=" + f, l["unite"], l["expr"]])
                ref[l["cle"]] = f"'Calcul'!$B${c.max_row}"
                if l["unite"] == "CHF":
                    c.cell(c.max_row, 2).number_format = "#,##0.00"
        s = wb.create_sheet("Sources")
        s.append(["Barème", "Valeur", "Source officielle", "Vérifié le"])
        for x in s[1]:
            x.font = gras
        for b in self.sources:
            s.append([f"{b['id']} {b['nom']}.{b['cle']}", b["valeur"], b["source"], b["verifie_le"]])
        for r in self.reserves:
            s.append(["Réserve", "", r, ""])
        for ws in (h, c, s):
            for col, w in zip("ABCD", (52, 18, 10, 90)):
                ws.column_dimensions[col].width = w
            for row in ws.iter_rows():
                for x in row:
                    x.alignment = Alignment(wrap_text=True, vertical="top")
        Path(chemin).parent.mkdir(parents=True, exist_ok=True)
        wb.save(chemin)
        return str(chemin)


def sortir(calc, excel=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # Windows
    except Exception:
        pass
    out = calc.resultat()
    if excel:
        out["excel"] = calc.excel(excel)
    print(json.dumps(out, ensure_ascii=False, indent=1))
    return out

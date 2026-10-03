"""Gabarits de la maison (§7.3), générés depuis le système de design : régénérables dès que la charte change.
python gabarits.py                 → (ré)génère Bureau/Modèles/* (idempotent)
python gabarits.py --inscrire      → en plus, inscrit/actualise chaque gabarit dans cerebro (GAB-…)
Sortie JSON : {gabarit: chemin}. Appelé aussi par produire.py quand un gabarit manque (racine neuve)."""
import sys, json, argparse, io
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from commun import MODELES, ROOT, journal, today
import design as D
import docx_outils as W

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GABARITS = {
    "memo.docx": "Avis de droit / mémo : page de titre, résumé exécutif, question, faits, droit sourcé, analyse, options, risques, recommandation, réserves, niveau de confort, annexe des sources ; en-tête Confidentiel, pied version + état du droit, pagination, titres numérotés",
    "lettre.docx": "Lettre de la maison : en-tête raison sociale et adresse, destinataire, lieu et date, objet, corps, formule, signature",
    "pv-assemblee.docx": "Procès-verbal d'assemblée (AG/associés/CA) : ouverture, présences, quorum, ordre du jour, décisions, clôture, signatures",
    "modele-calcul.xlsx": "Modèle de calcul : onglets Hypothèses / Calculs / Sensibilités / Sources ; noms définis, aucune valeur en dur dans les formules",
    "presentation.pptx": "Présentation 16:9 : thème de la maison, titres-affirmations, un message par diapositive, notes de l'orateur",
    "gabarit-rapport.pdf": "Gabarit PDF (reportlab, rendu_pdf.py) : page de titre, en-tête Confidentiel, pied version + état du droit, pagination ; spécimen des styles",
}


# ------------------------------------------------------------------ Word
def _base_docx(d, langue="fr"):
    doc = Document()
    W.mise_en_page(doc, d)
    W.configurer_styles(doc, d, langue)
    W.numerotation_titres(doc, d)
    core = doc.core_properties
    core.author = d["identite"]["raison_sociale"]; core.comments = f"Gabarit de la maison, design v{d['version']}"
    return doc


def _guide(doc, txt):
    p = doc.add_paragraph(txt, style="Maison Petit")
    return p


def memo(d):
    doc = _base_docx(d)
    W.entete_pied(doc, d)
    doc.add_paragraph("{{raison_sociale}}", style="Maison Petit")
    for _ in range(5):
        doc.add_paragraph()
    doc.add_paragraph("{{titre}}", style="Title")
    doc.add_paragraph("{{sous_titre}}", style="Subtitle")
    doc.add_paragraph()
    t = W.tableau(doc, d, ["Destinataire", "{{client}}"], [["Date", "{{date}}"], ["État du droit au", "{{date_etat}}"], ["Niveau de confort", "{{confort}}"], ["Version", "{{version}}"], ["Auteur", "{{auteur}}"]])
    W.saut_de_page(doc)
    sections = [("Résumé exécutif", "Une page au plus : conclusion d'abord, chiffres clés, recommandation et niveau de confort."),
                ("Question", "La question posée, en une ou deux phrases ; termes définis."),
                ("Faits", "Faits retenus, datés et étiquetés ([fait vérifié], [déclaré par … le …], [hypothèse])."),
                ("Droit applicable", "Chaque règle porte sa source datée (identifiant BIB-… ou référence officielle avec état du droit) ; sinon ⚠."),
                ("Analyse", "Une idée par paragraphe ; position assumée ; « so what » explicite."),
                ("Options", "Options chiffrées, avec avantages, inconvénients, coûts et délais."),
                ("Risques", "Risques, probabilité, impact, parade."),
                ("Recommandation", "La voie recommandée et la prochaine étape datée."),
                ("Réserves", "Limites de l'analyse, faits non vérifiés, contrôles non aboutis."),
                ("Niveau de confort", "Élevé / moyen / faible, avec la raison.")]
    for titre, aide in sections:
        doc.add_paragraph(titre, style="Heading 1")
        _guide(doc, aide)
    h = doc.add_paragraph("Annexe — Sources", style="Heading 1"); W.sans_numero(h)
    W.tableau(doc, d, ["Réf.", "Source", "Version / état", "Lien"], [["[1]", "{{source}}", "{{etat}}", "{{url}}"]])
    return doc


def lettre(d):
    doc = _base_docx(d)
    W.entete_pied(doc, d, gauche_h="{{raison_sociale}}", droite_h="{{confidentialite}}", gauche_p="{{adresse}}")
    doc.add_paragraph("{{destinataire}}", style="Maison Destinataire")
    p = doc.add_paragraph("{{lieu}}, le {{date}}"); p.paragraph_format.space_before = Pt(24); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph("{{objet}}", style="Maison Objet")
    doc.add_paragraph("{{salutation}}")
    doc.add_paragraph("{{corps}}")
    doc.add_paragraph("{{formule}}")
    doc.add_paragraph("{{signature}}", style="Maison Signature")
    doc.add_paragraph("{{annexes}}", style="Maison Petit")
    return doc


def pv(d):
    doc = _base_docx(d)
    W.entete_pied(doc, d)
    doc.add_paragraph("Procès-verbal", style="Title")
    doc.add_paragraph("{{titre}}", style="Subtitle")
    W.tableau(doc, d, ["Société", "{{societe}}"], [["Date et lieu", "{{date}} · {{lieu}}"], ["Présidence", "{{president}}"], ["Procès-verbal", "{{secretaire}}"]])
    for titre, aide in [("Ouverture", "Heure d'ouverture, présidence, désignation du secrétaire et des scrutateurs."),
                        ("Présences et représentation", "Actions/parts représentées, procurations, représentant indépendant le cas échéant."),
                        ("Constitution et quorum", "Convocation (forme, délai, source statutaire) ; constat de la constitution régulière."),
                        ("Ordre du jour", "Points inscrits à l'ordre du jour."),
                        ("Décisions", "Pour chaque point : proposition, discussion, vote (pour / contre / abstentions), décision."),
                        ("Clôture", "Heure de clôture ; absence d'opposition au procès-verbal.")]:
        doc.add_paragraph(titre, style="Heading 1")
        _guide(doc, aide)
    doc.add_paragraph("Signatures", style="Heading 1")
    W.tableau(doc, d, ["Le président", "Le secrétaire"], [["\n\n………………………………", "\n\n………………………………"], ["{{president}}", "{{secretaire}}"]])
    return doc


# ------------------------------------------------------------------ Excel
def calcul(d):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
    from openpyxl.workbook.defined_name import DefinedName
    wb = Workbook()
    hx = lambda n: D.couleur(d, n).lstrip("#")
    corps = d["typographie"]["corps"]["police"]
    st_ent = NamedStyle(name="Maison En-tête", font=Font(name=corps, bold=True, color=hx("blanc"), size=10), fill=PatternFill("solid", fgColor=hx("primaire")), alignment=Alignment(vertical="center", wrap_text=True))
    st_saisie = NamedStyle(name="Maison Saisie", font=Font(name=corps, color=hx("accent"), size=10), fill=PatternFill("solid", fgColor=hx("fond_clair")), border=Border(bottom=Side(style="thin", color=hx("filet"))))
    st_calc = NamedStyle(name="Maison Calcul", font=Font(name=corps, color=hx("encre"), size=10), border=Border(bottom=Side(style="thin", color=hx("filet"))))
    st_txt = NamedStyle(name="Maison Texte", font=Font(name=corps, color=hx("encre"), size=10), alignment=Alignment(wrap_text=True, vertical="top"))
    st_titre = NamedStyle(name="Maison Titre", font=Font(name=d["typographie"]["titres"]["police"], color=hx("primaire"), size=14, bold=True))
    for s in (st_ent, st_saisie, st_calc, st_txt, st_titre):
        wb.add_named_style(s)
    ws = wb.active; ws.title = "Hypothèses"
    ws["A1"] = "Hypothèses"; ws["A1"].style = "Maison Titre"
    ws["A2"] = "Saisir ici toute valeur (cellules ocre) ; chaque valeur porte sa source. Les formules ne contiennent aucune valeur en dur."; ws["A2"].style = "Maison Texte"
    ent = ["Nom", "Libellé", "Valeur", "Unité", "Source (BIB-… ou référence datée)", "Vérifié le"]
    for i, h in enumerate(ent, 1):
        c = ws.cell(row=4, column=i, value=h); c.style = "Maison En-tête"
    exemples = [("montant_base", "Montant de base", None, "CHF", "", ""), ("taux", "Taux applicable", None, "%", "⚠ à sourcer", ""), ("variation", "Pas de sensibilité", None, "%", "hypothèse de travail", "")]
    for r, (nom, lib, val, u, src, ver) in enumerate(exemples, 5):
        for i, v in enumerate((nom, lib, val, u, src, ver), 1):
            c = ws.cell(row=r, column=i, value=v); c.style = "Maison Saisie" if i == 3 else "Maison Texte"
        wb.defined_names[nom] = DefinedName(nom, attr_text=f"'Hypothèses'!$C${r}")
    for col, w in zip("ABCDEF", (18, 34, 14, 8, 40, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wc = wb.create_sheet("Calculs")
    wc["A1"] = "Calculs"; wc["A1"].style = "Maison Titre"
    for i, h in enumerate(["Étape", "Libellé", "Résultat", "Unité", "Formule (lisible)", "Note"], 1):
        c = wc.cell(row=3, column=i, value=h); c.style = "Maison En-tête"
    lignes = [("1", "Montant soumis", "=montant_base", "CHF", "montant_base", ""), ("2", "Résultat", "=C4*taux/100", "CHF", "étape 1 × taux", "")]
    for r, ligne in enumerate(lignes, 4):
        for i, v in enumerate(ligne, 1):
            c = wc.cell(row=r, column=i, value=v); c.style = "Maison Calcul" if i == 3 else "Maison Texte"
    for col, w in zip("ABCDEF", (8, 34, 16, 8, 34, 30)):
        wc.column_dimensions[col].width = w
    wsn = wb.create_sheet("Sensibilités")
    wsn["A1"] = "Sensibilités"; wsn["A1"].style = "Maison Titre"
    for i, h in enumerate(["Écart sur le taux (pas)", "Taux testé", "Résultat"], 1):
        c = wsn.cell(row=3, column=i, value=h); c.style = "Maison En-tête"
    for k, r in enumerate(range(4, 9)):
        wsn.cell(row=r, column=1, value=f"={k - 2}*variation").style = "Maison Calcul"
        wsn.cell(row=r, column=2, value=f"=taux+A{r}").style = "Maison Calcul"
        wsn.cell(row=r, column=3, value=f"=Calculs!$C$4*B{r}/100").style = "Maison Calcul"
    for col, w in zip("ABC", (24, 14, 16)):
        wsn.column_dimensions[col].width = w
    wsr = wb.create_sheet("Sources")
    wsr["A1"] = "Sources"; wsr["A1"].style = "Maison Titre"
    for i, h in enumerate(["Réf.", "Source", "Version / état du droit", "Lien", "Vérifié le"], 1):
        c = wsr.cell(row=3, column=i, value=h); c.style = "Maison En-tête"
    for col, w in zip("ABCDE", (10, 50, 22, 40, 12)):
        wsr.column_dimensions[col].width = w
    return wb


# ------------------------------------------------------------------ PowerPoint
def _theme_pptx(prs, d):
    """couleurs et polices de la maison dans le thème (titres, corps, accents)"""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    cs = root.find(".//a:clrScheme", ns)
    pal = d["palette"]
    mapping = {"dk1": D.couleur(d, "encre"), "lt1": "#FFFFFF", "dk2": D.couleur(d, "primaire"), "lt2": D.couleur(d, "fond_clair"),
               "accent1": pal["graphiques"][0], "accent2": pal["graphiques"][1], "accent3": pal["graphiques"][2],
               "accent4": pal["graphiques"][3], "accent5": pal["graphiques"][4], "accent6": pal["graphiques"][5],
               "hlink": D.couleur(d, "primaire"), "folHlink": D.couleur(d, "secondaire")}
    if cs is not None:
        cs.set("name", "Maison")
        for tag, hexa in mapping.items():
            el = cs.find(f"a:{tag}", ns)
            if el is None:
                continue
            for ch in list(el):
                el.remove(ch)
            s = etree.SubElement(el, f"{{{ns['a']}}}srgbClr"); s.set("val", hexa.lstrip("#").upper())
    fs = root.find(".//a:fontScheme", ns)
    if fs is not None:
        fs.set("name", "Maison")
        for tag, police in (("majorFont", d["typographie"]["titres"]["police"]), ("minorFont", d["typographie"]["corps"]["police"])):
            lat = fs.find(f"a:{tag}/a:latin", ns)
            if lat is not None:
                lat.set("typeface", police)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def presentation(d):
    from pptx import Presentation
    from pptx.util import Cm, Pt as PPt
    from pptx.dml.color import RGBColor as PRGB
    prs = Presentation()
    g = d["grille"]["presentation"]
    prs.slide_width, prs.slide_height = Cm(g["largeur_cm"]), Cm(g["hauteur_cm"])
    _theme_pptx(prs, d)
    # repositionner les zones des dispositions sur le format 16:9
    for layout in prs.slide_layouts:
        for ph in layout.placeholders:
            try:
                ph.left = int(ph.left * g["largeur_cm"] / 25.4)
                ph.width = int(ph.width * g["largeur_cm"] / 25.4)
            except Exception:
                pass
    for ph in prs.slide_master.placeholders:
        try:
            ph.left = int(ph.left * g["largeur_cm"] / 25.4); ph.width = int(ph.width * g["largeur_cm"] / 25.4)
        except Exception:
            pass
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "{{titre}}"
    s.placeholders[1].text = "{{client}} · {{date}}"
    s.notes_slide.notes_text_frame.text = "Notes de l'orateur : le message de la présentation en une phrase."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Le titre affirme le message de la diapositive"
    tf = s2.placeholders[1].text_frame
    tf.text = "Un seul message par diapositive"
    for t in ("Trois à cinq éléments au plus, chacun prouve le titre", "Chiffres avec unité et source"):
        p = tf.add_paragraph(); p.text = t
    s2.notes_slide.notes_text_frame.text = "Notes : ce que l'orateur dit, sources détaillées."
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = PRGB(*D.rgb(d, "primaire" if sh == sl.shapes.title else "encre"))
    return prs


# ------------------------------------------------------------------ génération
def generer(dest=None, inscrire=False):
    d = D.charger(force=True)
    dest = Path(dest or MODELES)
    dest.mkdir(parents=True, exist_ok=True)
    res = {}
    for nom, fn in (("memo.docx", memo), ("lettre.docx", lettre), ("pv-assemblee.docx", pv)):
        doc = fn(d)
        doc.save(dest / nom); res[nom] = str(dest / nom)
    wb = calcul(d); wb.save(dest / "modele-calcul.xlsx"); res["modele-calcul.xlsx"] = str(dest / "modele-calcul.xlsx")
    prs = presentation(d); prs.save(dest / "presentation.pptx"); res["presentation.pptx"] = str(dest / "presentation.pptx")
    try:
        import rendu_pdf
        rendu_pdf.specimen(dest / "gabarit-rapport.pdf", d); res["gabarit-rapport.pdf"] = str(dest / "gabarit-rapport.pdf")
    except Exception as e:
        journal("erreurs-producteur", op="gabarit-pdf", erreur=repr(e))
    if inscrire:
        res["_ids"] = inscrire_cerebro(res)
    journal("producteur", op="gabarits", n=len(res), design=d["version"])
    return res


def inscrire_cerebro(res):
    """un objet GAB-… par gabarit (créé une fois, puis actualisé)"""
    from commun import cb
    core, objets, _, _ = cb()
    con = core.db()
    ids = {}
    for nom, chemin in res.items():
        if nom.startswith("_"):
            continue
        rel = objets.relpath(chemin)
        r = con.execute("SELECT id FROM objets WHERE type='gabarit' AND nom=?", (f"gabarit {nom}",)).fetchone()
        if r:
            objets.update(r[0], chemin=rel, resume=GABARITS[nom], source=".equipe/scripts/producteur/gabarits.py")
            ids[nom] = r[0]
        else:
            gid = objets.create("gabarit", f"gabarit {nom}", resume=GABARITS[nom], source=".equipe/scripts/producteur/gabarits.py",
                                chemin=rel, prochaine_action="régénérer si la charte change (python gabarits.py)",
                                mots_cles="gabarit modèle design " + nom.split(".")[0])
            ids[nom] = gid
    return ids


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dest")
    ap.add_argument("--inscrire", action="store_true")
    a = ap.parse_args()
    print(json.dumps(generer(a.dest, a.inscrire), ensure_ascii=False))

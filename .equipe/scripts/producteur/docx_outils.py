"""Outils Word (python-docx) partagés par les gabarits et le rendu : styles de la maison, titres numérotés,
en-têtes/pieds avec champs, tableaux de la maison, table des matières."""
import copy
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor, Emu

import design as D

STYLES_MAISON = ["Maison Encadré", "Maison Réserve", "Maison Petit", "Maison Signature", "Maison Destinataire", "Maison Objet"]
NUM_ID_TITRES = 90  # identifiant réservé à la numérotation des titres de la maison


def _rgb(d, nom):
    return RGBColor(*D.rgb(d, nom))


def _police(style_or_run, nom):
    """applique la police aussi aux scripts est-asiatiques/complexes (sinon Word garde la police du thème)"""
    f = style_or_run.font
    f.name = nom
    el = style_or_run.element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), nom)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]


def _style(doc, nom, typ=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    try:
        return doc.styles[nom]
    except KeyError:
        s = doc.styles.add_style(nom, typ)
        if base and typ == WD_STYLE_TYPE.PARAGRAPH:
            s.base_style = doc.styles[base]
        return s


def configurer_styles(doc, d, langue="fr"):
    ty = d["typographie"]
    corps, titres = ty["corps"], ty["titres"]
    n = doc.styles["Normal"]
    _police(n, corps["police"])
    n.font.size = Pt(corps["taille"])
    n.font.color.rgb = _rgb(d, "encre")
    pf = n.paragraph_format
    pf.space_after = Pt(corps["apres"])
    pf.line_spacing = corps["interligne"]
    pf.widow_control = True
    # langue du document (correcteur Word, césure)
    lang = {"fr": "fr-CH", "de": "de-CH", "it": "it-CH", "en": "en-GB"}.get(langue, "fr-CH")
    rpr = n.element.get_or_add_rPr()
    l = rpr.find(qn("w:lang"))
    if l is None:
        l = OxmlElement("w:lang"); rpr.append(l)
    l.set(qn("w:val"), lang)
    for niv in ty["niveaux"]:
        s = doc.styles[niv["style"]]
        _police(s, titres["police"])
        s.font.size = Pt(niv["taille"])
        s.font.bold = niv.get("gras", True)
        s.font.italic = niv.get("italique", False)
        s.font.color.rgb = _rgb(d, titres["couleur"])
        s.paragraph_format.space_before = Pt(niv["avant"])
        s.paragraph_format.space_after = Pt(niv["apres"])
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.keep_together = True
    t = doc.styles["Title"]
    _police(t, titres["police"]); t.font.size = Pt(ty["titre_document"]["taille"]); t.font.color.rgb = _rgb(d, ty["titre_document"]["couleur"]); t.font.bold = False
    _bordure_paragraphe(t, None)
    st = doc.styles["Subtitle"]
    _police(st, corps["police"]); st.font.size = Pt(ty["sous_titre"]["taille"]); st.font.color.rgb = _rgb(d, ty["sous_titre"]["couleur"]); st.font.italic = False
    cap = doc.styles["Caption"]
    _police(cap, corps["police"]); cap.font.size = Pt(d["legendes"]["taille"]); cap.font.color.rgb = _rgb(d, d["legendes"]["couleur"]); cap.font.bold = False; cap.font.italic = True
    for nom in ("List Bullet", "List Number", "Quote"):
        s = doc.styles[nom]; _police(s, corps["police"]); s.font.size = Pt(corps["taille"])
    q = doc.styles["Quote"]; q.font.color.rgb = _rgb(d, "secondaire"); q.paragraph_format.left_indent = Mm(8)
    # styles de la maison
    e = _style(doc, "Maison Encadré"); e.paragraph_format.left_indent = Mm(4); e.paragraph_format.right_indent = Mm(4)
    _ombrage_style(e, D.couleur(d, "fond_clair")); _bordure_gauche(e, D.couleur(d, "primaire"))
    r = _style(doc, "Maison Réserve"); r.font.color.rgb = _rgb(d, "alerte"); r.font.size = Pt(corps["taille"] - 0.5)
    p = _style(doc, "Maison Petit"); p.font.size = Pt(9); p.font.color.rgb = _rgb(d, "discret")
    sg = _style(doc, "Maison Signature"); sg.paragraph_format.space_before = Pt(24); sg.paragraph_format.keep_together = True
    ds = _style(doc, "Maison Destinataire"); ds.paragraph_format.space_after = Pt(0); ds.paragraph_format.left_indent = Mm(95)
    ob = _style(doc, "Maison Objet"); ob.font.bold = True; ob.paragraph_format.space_before = Pt(18); ob.paragraph_format.space_after = Pt(12)
    for nom in ("Header", "Footer"):
        s = doc.styles[nom]; _police(s, corps["police"]); s.font.size = Pt(d["entete_pied"]["taille"]); s.font.color.rgb = _rgb(d, d["entete_pied"]["couleur"])
        ppr = s.element.get_or_add_pPr()
        for tabs in ppr.findall(qn("w:tabs")):  # tabulations centre/droite du modèle par défaut (Letter) : retirées
            ppr.remove(tabs)
    for i in (1, 2, 3):
        try:
            s = doc.styles[f"TOC {i}"]
        except KeyError:
            s = doc.styles.add_style(f"TOC {i}", WD_STYLE_TYPE.PARAGRAPH); s.base_style = doc.styles["Normal"]
        s.paragraph_format.left_indent = Mm(5 * (i - 1)); s.paragraph_format.space_after = Pt(2)
        s.paragraph_format.tab_stops.add_tab_stop(Mm(160), WD_TAB_ALIGNMENT.RIGHT)


def _ombrage_style(style, hexa):
    ppr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexa.lstrip("#"))
    ppr.append(shd)


def _bordure_gauche(style, hexa):
    ppr = style.element.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "6"), ("color", hexa.lstrip("#"))):
        l.set(qn(f"w:{k}"), v)
    b.append(l); ppr.append(b)


def _bordure_paragraphe(style, hexa):
    ppr = style.element.get_or_add_pPr()
    for b in ppr.findall(qn("w:pBdr")):
        ppr.remove(b)


def numerotation_titres(doc, d):
    """liste multiniveau liée aux styles Heading 1..3 (1. / 1.1 / 1.1.1)"""
    numbering = doc.part.numbering_part.element
    for an in numbering.findall(qn("w:abstractNum")):
        if an.get(qn("w:abstractNumId")) == str(NUM_ID_TITRES):
            return
    an = OxmlElement("w:abstractNum"); an.set(qn("w:abstractNumId"), str(NUM_ID_TITRES))
    mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "multilevel"); an.append(mlt)
    for i, niv in enumerate(d["typographie"]["niveaux"]):
        lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), str(i))
        for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:pStyle", niv["style"].replace(" ", "")), ("w:lvlText", niv["numero"]), ("w:lvlJc", "left")):
            e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
        ppr = OxmlElement("w:pPr"); ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(int(567 + 284 * i))); ind.set(qn("w:hanging"), str(int(567 + 284 * i))); ppr.append(ind); lvl.append(ppr)
        an.append(lvl)
    # abstractNum avant les w:num (ordre imposé par le schéma)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(an)
    else:
        numbering.append(an)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(NUM_ID_TITRES))
    a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), str(NUM_ID_TITRES)); num.append(a)
    numbering.append(num)
    for i, niv in enumerate(d["typographie"]["niveaux"]):
        ppr = doc.styles[niv["style"]].element.get_or_add_pPr()
        for old in ppr.findall(qn("w:numPr")):
            ppr.remove(old)
        numpr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(i)); numpr.append(il)
        ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(NUM_ID_TITRES)); numpr.append(ni)
        ppr.insert(0, numpr)


def sans_numero(paragraph):
    """titre non numéroté (annexes, page de titre)"""
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "0"); numpr.append(ni)
    ppr.insert(0, numpr)


def mise_en_page(doc, d):
    g = d["grille"]["marges_mm"]
    for s in doc.sections:
        s.page_height, s.page_width = Mm(297), Mm(210)
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Mm(g["haut"]), Mm(g["bas"]), Mm(g["gauche"]), Mm(g["droite"])
        s.header_distance, s.footer_distance = Mm(d["grille"]["entete_mm"]), Mm(d["grille"]["pied_mm"])


def largeur_texte(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def champ(paragraph, instr, placeholder="1"):
    """champ Word (PAGE, NUMPAGES, TOC …) avec résultat en cache"""
    r = paragraph.add_run()
    for typ in ("begin",):
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); r._r.append(fc)
    r2 = paragraph.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r2._r.append(it)
    r3 = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r3._r.append(fc)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r5._r.append(fc)
    return r4


def _tabs_droite(paragraph, doc):
    paragraph.paragraph_format.tab_stops.add_tab_stop(largeur_texte(doc), WD_TAB_ALIGNMENT.RIGHT)


def entete_pied(doc, d, gauche_h="{{raison_sociale}}", droite_h="{{confidentialite}}", gauche_p="{{pied}}", pages=True):
    """en-tête : raison sociale | confidentialité ; pied : version et date d'état du droit | Page x / y.
    Les {{…}} sont remplacés à la production (remplacer_marqueurs)."""
    for s in doc.sections:
        h = s.header; h.is_linked_to_previous = False
        p = h.paragraphs[0]; p.text = ""; p.style = doc.styles["Header"]; _tabs_droite(p, doc)
        p.add_run(gauche_h); p.add_run("\t")
        rr = p.add_run(droite_h); rr.bold = True; rr.font.color.rgb = _rgb(d, "primaire")
        _filet_bas(p, D.couleur(d, "filet"))
        f = s.footer; f.is_linked_to_previous = False
        q = f.paragraphs[0]; q.text = ""; q.style = doc.styles["Footer"]; _tabs_droite(q, doc)
        q.add_run(gauche_p); q.add_run("\t")
        if pages:
            q.add_run("Page "); champ(q, "PAGE"); q.add_run(" / "); champ(q, "NUMPAGES")


def _filet_bas(paragraph, hexa):
    ppr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); l = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", hexa.lstrip("#"))):
        l.set(qn(f"w:{k}"), v)
    b.append(l); ppr.append(b)


def remplacer_marqueurs(doc, valeurs):
    """remplace {{clé}} dans corps, en-têtes et pieds (marqueurs écrits dans un seul run par les gabarits)"""
    def _par(p):
        for r in p.runs:
            if "{{" in r.text:
                t = r.text
                for k, v in valeurs.items():
                    t = t.replace("{{" + k + "}}", str(v))
                r.text = t
    parts = [doc]
    for s in doc.sections:
        parts += [s.header, s.footer, s.first_page_header, s.first_page_footer]
    for part in parts:
        for p in part.paragraphs:
            _par(p)
        for t in part.tables:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        _par(p)


def ombrer_cellule(cell, hexa):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexa.lstrip("#"))
    tcpr.append(shd)


def tableau(doc, d, entete, lignes, ajouter_runs=None, largeurs=None):
    """tableau de la maison : en-tête foncé répété, lignes alternées, lignes insécables, largeur = zone de texte"""
    t = doc.add_table(rows=1, cols=len(entete))
    try:
        t.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb = d["tableaux"]
    W = largeur_texte(doc)
    n = len(entete)
    larg = largeurs or [int(W / n)] * n
    def _remplir(cell, txt, gras=False, couleur=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        if ajouter_runs:
            ajouter_runs(p, txt, gras=gras, taille=tb["taille"], couleur=couleur)
        else:
            r = p.add_run(txt); r.bold = gras; r.font.size = Pt(tb["taille"])
            if couleur:
                r.font.color.rgb = couleur
    for i, h in enumerate(entete):
        c = t.rows[0].cells[i]
        _remplir(c, h, gras=True, couleur=_rgb(d, tb["entete"]["texte"]))
        ombrer_cellule(c, D.couleur(d, tb["entete"]["fond"]))
    trpr = t.rows[0]._tr.get_or_add_trPr()
    if tb.get("repeter_entete"):
        e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trpr.append(e)
    for k, ligne in enumerate(lignes):
        row = t.add_row()
        for i in range(n):
            txt = ligne[i] if i < len(ligne) else ""
            _remplir(row.cells[i], txt)
            if k % 2 == 1:
                ombrer_cellule(row.cells[i], D.couleur(d, tb["lignes_alternees"]))
        if tb.get("lignes_insecables"):
            e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(e)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Emu(larg[i])
    return t


def saut_de_page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def table_des_matieres(doc, d, titres, apres_paragraphe=None, libelle="Table des matières"):
    """champ TOC avec résultat en cache (titres + pages connues) ; Word le met à jour d'un clic droit."""
    paras = []
    h = doc.add_paragraph(libelle, style="Heading 1"); sans_numero(h); paras.append(h)
    p = doc.add_paragraph(style="TOC 1")
    r = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r2 = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-3" \\h \\z \\u '; r2._r.append(it)
    r3 = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r3._r.append(fc)
    paras.append(p)
    first = True
    for (niv, num, txt, page) in titres:
        q = p if first else doc.add_paragraph(style=f"TOC {min(niv, 3)}")
        if first:
            q.style = doc.styles[f"TOC {min(niv, 3)}"]
        first = False
        q.add_run(f"{num}  {txt}\t{page if page else ''}")
        if q is not p:
            paras.append(q)
    end = doc.add_paragraph()
    r5 = end.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r5._r.append(fc)
    paras.append(end)
    if apres_paragraphe is not None:  # déplacer le bloc juste après un paragraphe donné
        anchor = apres_paragraphe._p
        for q in paras:
            anchor.addnext(q._p); anchor = q._p
    return paras


def vider_corps(doc):
    """supprime le contenu d'exemple du gabarit en gardant styles, sections, en-têtes et pieds"""
    body = doc.element.body
    for el in list(body):
        if el.tag == qn("w:sectPr"):
            continue
        body.remove(el)


def styles_presents(doc):
    return {s.name for s in doc.styles}

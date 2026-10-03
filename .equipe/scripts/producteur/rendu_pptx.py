"""Rendu PowerPoint depuis presentation.pptx : diapositive de titre, puis une diapositive par titre de niveau 1
(titre-affirmation, puces, image ou tableau, notes de l'orateur depuis un bloc ```notes``` ou « Notes : »)."""
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor

import design as D
import mdparse as MD
from commun import MODELES, date_longue


def _supprimer_diapos(prs):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        rid = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        prs.part.drop_rel(rid)
        ids.remove(sld)


def _couleur(r, d, nom):
    r.font.color.rgb = RGBColor(*D.rgb(d, nom))


def rendre(meta, blocs, dest, d):
    p = MODELES / "presentation.pptx"
    if not p.exists():
        import gabarits
        gabarits.generer(MODELES)
    prs = Presentation(str(p))
    _supprimer_diapos(prs)
    lg = meta.get("langue", "fr")
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = MD.texte_brut(meta.get("titre", ""))
    for para in s.shapes.title.text_frame.paragraphs:
        for r in para.runs:
            r.font.size = Pt(34); _couleur(r, d, "primaire")
    s.placeholders[1].text = f"{meta.get('client_nom', '')} · {date_longue(meta.get('date'), lg)}"
    s.notes_slide.notes_text_frame.text = MD.texte_brut(meta.get("message", meta.get("titre", "")))
    W = prs.slide_width; H = prs.slide_height
    for titre, bl in MD.sections(blocs, 1):
        if titre is None:
            continue
        sl = prs.slides.add_slide(prs.slide_layouts[1] if not any(b["t"] in ("image", "table") for b in bl) else prs.slide_layouts[5])
        sl.shapes.title.text = MD.texte_brut(titre)
        for para in sl.shapes.title.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(28); _couleur(r, d, "primaire")
        puces, notes = [], []
        for b in bl:
            if b["t"] in ("ul", "ol"):
                puces += [MD.texte_brut(x) for x in b["items"]]
            elif b["t"] == "p":
                t = MD.texte_brut(b["texte"])
                if re.match(r"(?i)^(notes?|notizen|note dell'oratore)\s*:", t):
                    notes.append(t.split(":", 1)[1].strip())
                else:
                    puces.append(t)
            elif b["t"] == "code" and b["lang"] in ("notes", "note"):
                notes.append(b["texte"])
            elif b["t"] == "image" and Path(b["chemin"]).exists():
                sl.shapes.add_picture(b["chemin"], Cm(2), Cm(4.2), height=H - Cm(6))
            elif b["t"] == "table":
                rows, cols = len(b["lignes"]) + 1, len(b["entete"])
                tb = sl.shapes.add_table(rows, cols, Cm(2), Cm(4.5), W - Cm(4), Cm(1) * rows).table
                for j, h in enumerate(b["entete"]):
                    tb.cell(0, j).text = MD.texte_brut(h)
                for i, r in enumerate(b["lignes"], 1):
                    for j in range(cols):
                        tb.cell(i, j).text = MD.texte_brut(r[j] if j < len(r) else "")
        if len(sl.placeholders) > 1 and puces:
            tf = sl.placeholders[1].text_frame
            tf.text = puces[0]
            for x in puces[1:]:
                tf.add_paragraph().text = x
            for para in tf.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(20); _couleur(r, d, "encre")
        elif puces:
            tx = sl.shapes.add_textbox(W - Cm(12), Cm(4.5), Cm(10), H - Cm(7)).text_frame
            tx.word_wrap = True
            tx.text = puces[0]
            for x in puces[1:]:
                tx.add_paragraph().text = x
        sl.notes_slide.notes_text_frame.text = "\n".join(notes) or ""
    prs.core_properties.title = meta.get("titre", "")
    prs.core_properties.category = "presentation"
    prs.core_properties.language = lg
    prs.save(str(dest))

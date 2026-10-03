"""Porte (f) présentation (§7.1, §7.3) : le livrable sort d'un gabarit de la maison et suit la structure attendue
pour son type (mémo, PV, lettre, calcul, présentation, mail)."""
import re
from pathlib import Path
from gates_common import resultat, ETAT_OK, ETAT_KO, ETAT_NA, design, section_alias
import common as C


def _structure(doc, d, typ):
    attendues = ((d.get("livrables") or {}).get(typ) or {}).get("sections") or []
    if not attendues:
        return [], []
    vues = {section_alias(t["texte"], d) for t in doc.get("titres", []) if t.get("niveau", 1) <= 2}
    manquantes = [s for s in attendues if s not in vues]
    return attendues, manquantes


def _docx(doc, d, typ, corr, det):
    from docx import Document
    obj = doc.get("objet")
    gab_nom = ((d.get("livrables") or {}).get(typ) or {}).get("gabarit") or "memo.docx"
    gab = C.MODELES / gab_nom
    styles_doc = {s.name for s in obj.styles}
    import docx_tools as W
    manquants = [s for s in W.STYLES_MAISON if s not in styles_doc]
    if manquants:
        corr.append({"probleme": f"styles de la maison absents : {', '.join(manquants)}", "correction": f"reproduire depuis Bureau/Modeles/{gab_nom} (produce.py)"})
    if gab.exists():
        styles_gab = {s.name for s in Document(str(gab)).styles}
        hors = sorted({p.style.name for p in obj.paragraphs if p.style is not None} - styles_gab)
        if hors:
            corr.append({"probleme": f"styles hors gabarit : {', '.join(hors[:8])}", "correction": "appliquer les styles du gabarit"})
    num = obj.part.numbering_part.element.xml if obj.part.package else ""
    if typ in ("memo", "note", "pv") and f'w:abstractNumId="{W.NUM_ID_TITRES}"' not in num:
        corr.append({"probleme": "titres non numérotés par le gabarit", "correction": "partir du gabarit (titres numérotés 1 / 1.1 / 1.1.1)"})
    entete = " ".join(p.text for s in obj.sections for p in s.header.paragraphs)
    pied = " ".join(p.text for s in obj.sections for p in s.footer.paragraphs)
    pied_xml = " ".join(s.footer._element.xml for s in obj.sections)
    if typ in ("memo", "note", "pv"):
        conf = d.get("identite", {}).get("mention_confidentialite", {}).values()
        if not any(c in entete for c in conf):
            corr.append({"probleme": "mention de confidentialité absente de l'en-tête", "correction": "en-tête du gabarit"})
        if not re.search(r"Version|Stand|Versione", pied) or not re.search(r"droit|Rechts|diritto|law", pied):
            corr.append({"probleme": "pied de page sans version ni date d'état du droit", "correction": "pied du gabarit"})
    if "PAGE" not in pied_xml:
        corr.append({"probleme": "pagination absente", "correction": "pied du gabarit (Page x / y)"})
    if "{{" in entete + pied + " ".join(p.text for p in obj.paragraphs):
        corr.append({"probleme": "marqueurs {{…}} non remplacés", "correction": "renseigner le front matter (titre, client, date_etat, confort)"})
    pages = (doc.get("ctx") or {}).get("pages")
    lim = ((d.get("livrables") or {}).get(typ) or {}).get("table_des_matieres_au_dela_pages")
    if pages and lim and pages > lim and "TOC" not in obj.element.body.xml:
        corr.append({"probleme": f"{pages} pages sans table des matières", "correction": "produce.py ajoute la table au-delà de 10 pages"})
    det.append(f"gabarit {gab_nom} · styles maison {'ok' if not manquants else 'incomplets'}")


def _xlsx(doc, d, corr, det):
    wb = doc["objet"]
    attendus = (d.get("livrables") or {}).get("calcul", {}).get("onglets", [])
    manq = [o for o in attendus if o not in wb.sheetnames]
    if manq:
        corr.append({"probleme": f"onglets manquants : {', '.join(manq)}", "correction": "partir de Bureau/Modeles/modele-calcul.xlsx"})
    durs = []
    for f in doc.get("formules", []):
        s = re.sub(r"'[^']*'!|\b[A-Za-z_][\w.]*!", "", f["formule"])          # noms de feuilles
        s = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", s)                              # références de cellules
        s = re.sub(r"\b[A-Za-z_][\w.]*\b", "", s)                              # noms définis, fonctions
        s = re.sub(r'"[^"]*"', "", s)
        lits = [x for x in re.findall(r"(?<![\w.])\d+(?:\.\d+)?", s) if x not in ("0", "1")]
        if lits:
            durs.append(f"{f['feuille']}!{f['cellule']} {f['formule']}")
    if durs:
        corr.append({"probleme": f"{len(durs)} formule(s) avec valeur en dur", "exemples": durs[:5], "correction": "déplacer la valeur dans l'onglet Hypothèses (nom défini) et la sourcer"})
    det.append(f"{len(doc.get('formules', []))} formules, {len(durs)} avec valeur en dur")


def _pptx(doc, d, corr, det):
    conf = (d.get("livrables") or {}).get("presentation", {})
    for s in doc.get("diapos", []):
        if s["n"] == 1:
            continue
        if not s["titre"]:
            corr.append({"diapo": s["n"], "probleme": "sans titre", "correction": "titre-affirmation"})
        elif len(s["titre"].split()) < conf.get("mots_titre_min", 4):
            corr.append({"diapo": s["n"], "probleme": f"titre « {s['titre']} » : étiquette, pas une affirmation", "correction": "écrire le message en phrase (« Le dividende peut être versé fin octobre »)"})
        if len(s["puces"]) > conf.get("puces_max", 5):
            corr.append({"diapo": s["n"], "probleme": f"{len(s['puces'])} puces (> {conf.get('puces_max', 5)}) : plus d'un message", "correction": "scinder la diapositive"})
        if not s["notes"]:
            corr.append({"diapo": s["n"], "probleme": "sans notes de l'orateur", "correction": "ajouter les notes"})
    det.append(f"{len(doc.get('diapos', []))} diapositives")


def _mail(doc, d, corr, det):
    m = doc.get("meta", {})
    if not (m.get("objet") or m.get("titre")):
        corr.append({"probleme": "objet du mail absent", "correction": "objet précis (sujet + décision attendue)"})
    sig = str(d.get("identite", {}).get("signature", "")).split(",")[0].strip()
    if sig and sig not in doc.get("texte", ""):
        corr.append({"probleme": "signature de la maison absente", "correction": "produce.py ajoute la signature"})
    for p in m.get("pieces") or []:
        n = Path(str(p)).name
        if not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*-\d{4}-\d{2}-\d{2}-v\d+\.[a-z0-9]+", n):
            corr.append({"probleme": f"pièce mal nommée : {n}", "correction": "client-objet-date-vN.ext"})
    det.append(f"mail · {len(m.get('pieces') or [])} pièce(s)")


def verifier(doc, ctx=None):
    d = design()
    typ = doc.get("type") or ""
    corr, det = [], []
    fmt = doc.get("format")
    try:
        if fmt == "docx" and doc.get("objet") is not None:
            _docx(doc, d, typ, corr, det)
        elif fmt == "xlsx" and doc.get("objet") is not None:
            _xlsx(doc, d, corr, det)
        elif fmt == "pptx":
            _pptx(doc, d, corr, det)
        elif typ == "mail" or fmt in ("eml",):
            _mail(doc, d, corr, det)
    except Exception as e:
        det.append(f"contrôle partiel ({e!r})")
    if typ and fmt in ("md", "docx", "pdf"):
        att, manq = _structure(doc, d, typ)
        if manq:
            corr.append({"probleme": f"structure {typ} incomplète : manque {', '.join(manq)}", "correction": "ajouter les sections (titres reconnus FR/DE/IT/EN, voir system.yaml alias_sections)"})
        if att:
            det.append(f"structure {typ} : {len(att) - len(manq)}/{len(att)} sections")
    if not det and not corr:
        return resultat(ETAT_NA, [f"pas de contrôle de présentation pour {fmt or '?'}"])
    return resultat(ETAT_KO if corr else ETAT_OK, det, corr)

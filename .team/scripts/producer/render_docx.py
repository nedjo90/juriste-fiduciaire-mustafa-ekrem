"""Rendu Word depuis les gabarits de la maison (python-docx) : mémo/avis/note, lettre, procès-verbal."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

import design as D
import docx_tools as W
import mdparse as MD
from common import MODELES, date_longue

LIB = {
    "fr": {"dest": "Destinataire", "date": "Date", "etat": "État du droit au", "confort": "Niveau de confort", "version": "Version", "auteur": "Auteur",
           "sources": "Annexe — Sources", "ref": "Réf.", "source": "Source", "vers": "Version / état", "lien": "Lien", "toc": "Table des matières",
           "pied": "Version {v} · État du droit au {d}", "le": "le", "annexes": "Annexes", "pv": "Procès-verbal", "societe": "Société",
           "datelieu": "Date et lieu", "pres": "Présidence", "secr": "Procès-verbal tenu par", "sig": "Signatures", "lepres": "Le président", "lesecr": "Le secrétaire"},
    "de": {"dest": "Empfänger", "date": "Datum", "etat": "Rechtsstand per", "confort": "Sicherheitsgrad", "version": "Version", "auteur": "Verfasser",
           "sources": "Anhang — Quellen", "ref": "Ref.", "source": "Quelle", "vers": "Fassung / Stand", "lien": "Link", "toc": "Inhaltsverzeichnis",
           "pied": "Version {v} · Rechtsstand per {d}", "le": "", "annexes": "Beilagen", "pv": "Protokoll", "societe": "Gesellschaft",
           "datelieu": "Datum und Ort", "pres": "Vorsitz", "secr": "Protokollführung", "sig": "Unterschriften", "lepres": "Der Vorsitzende", "lesecr": "Der Protokollführer"},
    "it": {"dest": "Destinatario", "date": "Data", "etat": "Stato del diritto al", "confort": "Grado di certezza", "version": "Versione", "auteur": "Autore",
           "sources": "Allegato — Fonti", "ref": "Rif.", "source": "Fonte", "vers": "Versione / stato", "lien": "Link", "toc": "Indice",
           "pied": "Versione {v} · Stato del diritto al {d}", "le": "", "annexes": "Allegati", "pv": "Verbale", "societe": "Società",
           "datelieu": "Data e luogo", "pres": "Presidenza", "secr": "Verbale redatto da", "sig": "Firme", "lepres": "Il presidente", "lesecr": "Il segretario"},
    "en": {"dest": "Addressee", "date": "Date", "etat": "Law as at", "confort": "Level of comfort", "version": "Version", "auteur": "Author",
           "sources": "Annex — Sources", "ref": "Ref.", "source": "Source", "vers": "Version / as at", "lien": "Link", "toc": "Contents",
           "pied": "Version {v} · Law as at {d}", "le": "", "annexes": "Enclosures", "pv": "Minutes", "societe": "Company",
           "datelieu": "Date and place", "pres": "Chair", "secr": "Minutes taken by", "sig": "Signatures", "lepres": "The chair", "lesecr": "The secretary"},
}


def lib(langue):
    return LIB.get(langue, LIB["fr"])


def ajouter_runs(d):
    alerte = W._rgb(d, "alerte")

    def _f(p, txt, gras=False, taille=None, couleur=None):
        for (t, b, i, code) in MD.runs(txt):
            for k, morceau in enumerate(re.split(r"(⚠)", t)):
                if not morceau:
                    continue
                r = p.add_run(morceau)
                r.bold = gras or b or (morceau == "⚠")
                r.italic = i
                if code:
                    r.font.name = d["typographie"]["mono"]["police"]
                if taille:
                    r.font.size = Pt(taille)
                if morceau == "⚠":
                    r.font.color.rgb = alerte
                elif couleur is not None:
                    r.font.color.rgb = couleur
    return _f


def ouvrir_gabarit(nom):
    p = MODELES / nom
    if not p.exists():
        import templates as gabarits
        gabarits.generer(MODELES)
    return Document(str(p))


def corps(doc, d, blocs, langue, sources_table=None):
    """blocs markdown → paragraphes aux styles du gabarit ; renvoie la liste des titres (niveau, numéro, texte)"""
    run = ajouter_runs(d)
    W_utile = W.largeur_texte(doc)
    titres, nums = [], [0, 0, 0]
    sources_vues = False
    attente = [False]  # table des sources à poser à la fin de la section « sources »
    n_fig = n_tab = 0

    def poser_sources():
        if attente[0] and sources_table:
            _table_sources(doc, d, sources_table, langue)
        attente[0] = False
    for b in blocs:
        t = b["t"]
        if t == "h" and b["niveau"] == 1:
            poser_sources()
        if t == "h":
            niv = min(b["niveau"], 3)
            from_alias = _alias(b["texte"], d)
            p = doc.add_paragraph(style=f"Heading {niv}")
            if from_alias == "sources" and niv == 1:
                run(p, lib(langue)["sources"] if b["texte"].lower().startswith(("annexe", "sources")) else MD.texte_brut(b["texte"]))
                W.sans_numero(p); sources_vues = True
                titres.append((1, "", MD.texte_brut(b["texte"]), None))
                attente[0] = True
                continue
            run(p, MD.texte_brut(b["texte"]))
            nums[niv - 1] += 1
            for k in range(niv, 3):
                nums[k] = 0
            titres.append((niv, ".".join(str(x) for x in nums[:niv]) + ("." if niv == 1 else ""), MD.texte_brut(b["texte"]), None))
        elif t == "p":
            p = doc.add_paragraph(); run(p, b["texte"])
            if b["texte"].strip().startswith("⚠"):
                p.style = doc.styles["Maison Réserve"]
        elif t == "quote":
            p = doc.add_paragraph(style="Quote"); run(p, b["texte"])
        elif t in ("ul", "ol"):
            for it in b["items"]:
                p = doc.add_paragraph(style="List Bullet" if t == "ul" else "List Number"); run(p, it)
        elif t == "table":
            n_tab += 1
            W.tableau(doc, d, [MD.texte_brut(x) for x in b["entete"]], [[c for c in r] for r in b["lignes"]], ajouter_runs=run)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "image" and Path(b["chemin"]).exists():
            n_fig += 1
            from PIL import Image
            with Image.open(b["chemin"]) as im:
                w_px, h_px = im.size
            larg = min(W_utile, Emu(int(w_px / 200 * 914400)))
            doc.add_picture(b["chemin"], width=larg)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            if b.get("legende"):
                cap = doc.add_paragraph(style="Caption"); run(cap, b["legende"])
        elif t == "saut":
            W.saut_de_page(doc)
    poser_sources()
    if sources_table and not sources_vues:
        p = doc.add_paragraph(lib(langue)["sources"], style="Heading 1"); W.sans_numero(p)
        titres.append((1, "", lib(langue)["sources"], None))
        _table_sources(doc, d, sources_table, langue)
    return titres


def _alias(titre, d):
    t = re.sub(r"^[\d.\s]+", "", MD.texte_brut(titre).strip().lower())
    for a in d["alias_sections"]["sources"]:
        if t == a or t.startswith(a) or t.startswith("annexe"):
            return "sources"
    return None


def _table_sources(doc, d, sources, langue):
    L = lib(langue)
    W.tableau(doc, d, [L["ref"], L["source"], L["vers"], L["lien"]],
              [[f"[{s['n']}]", s.get("titre", ""), s.get("etat", ""), s.get("url", "")] for s in sources], ajouter_runs=ajouter_runs(d),
              largeurs=[int(W.largeur_texte(doc) * x) for x in (0.08, 0.47, 0.2, 0.25)])


def _proprietes(doc, d, meta):
    cp = doc.core_properties
    cp.title = meta.get("titre", "")
    cp.category = meta.get("type", "")
    cp.language = meta.get("langue", "fr")
    cp.keywords = " ".join(x for x in [meta.get("client_id"), meta.get("dossier")] + list(meta.get("liens") or []) if x)
    cp.version = f"v{meta.get('version', 1)}"
    cp.author = d["identite"]["raison_sociale"]
    cp.subject = meta.get("objet", "")
    cp.comments = ""


def marqueurs(d, meta):
    L = lib(meta.get("langue", "fr"))
    lg = meta.get("langue", "fr")
    return {"raison_sociale": d["identite"]["raison_sociale"],
            "confidentialite": D.mention_confidentialite(d, lg) if meta.get("confidentiel", True) else "",
            "pied": L["pied"].format(v=meta.get("version", 1), d=date_longue(meta.get("date_etat") or meta.get("date"), lg)),
            "adresse": d["identite"].get("adresse") or d["identite"]["raison_sociale"]}


def memo(meta, blocs, dest, d, sources=None, titres_pages=None):
    """mémo / avis de droit / note : page de titre + corps + annexe des sources (+ table des matières si demandée)"""
    lg = meta.get("langue", "fr"); L = lib(lg)
    doc = ouvrir_gabarit("memo.docx")
    W.vider_corps(doc)
    run = ajouter_runs(d)
    doc.add_paragraph(d["identite"]["raison_sociale"], style="Maison Petit")
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title"); run(p, meta.get("titre", ""))
    if meta.get("sous_titre"):
        p = doc.add_paragraph(style="Subtitle"); run(p, meta["sous_titre"])
    doc.add_paragraph()
    W.tableau(doc, d, [L["dest"], meta.get("client_nom", "")],
              [[L["date"], date_longue(meta.get("date"), lg)], [L["etat"], date_longue(meta.get("date_etat"), lg)],
               [L["confort"], str(meta.get("confort", ""))], [L["version"], f"v{meta.get('version', 1)}"],
               [L["auteur"], meta.get("auteur") or d["identite"]["signature"]]], ajouter_runs=run)
    W.saut_de_page(doc)
    ancre = doc.paragraphs[-1]
    titres = corps(doc, d, blocs, lg, sources)
    if titres_pages is not None:
        W.table_des_matieres(doc, d, [(n, num, t, titres_pages.get((num, t))) for (n, num, t, _) in titres], apres_paragraphe=ancre, libelle=L["toc"])
    W.remplacer_marqueurs(doc, marqueurs(d, meta))
    _proprietes(doc, d, meta)
    doc.save(str(dest))
    return titres


def lettre(meta, blocs, dest, d, sources=None):
    lg = meta.get("langue", "fr"); L = lib(lg)
    doc = ouvrir_gabarit("lettre.docx")
    W.vider_corps(doc)
    run = ajouter_runs(d)
    for l in str(meta.get("destinataire", "")).strip().split("\n"):
        doc.add_paragraph(l.strip(), style="Maison Destinataire")
    lieu = meta.get("lieu") or ""
    p = doc.add_paragraph((f"{lieu}, {L['le']} " if lieu and L["le"] else (f"{lieu}, " if lieu else "")) + date_longue(meta.get("date"), lg))
    p.paragraph_format.space_before = Pt(24)
    p = doc.add_paragraph(style="Maison Objet"); run(p, meta.get("titre", ""))
    if meta.get("salutation"):
        doc.add_paragraph(meta["salutation"])
    corps(doc, d, [b for b in blocs if b["t"] != "h"], lg, None)
    if meta.get("formule"):
        doc.add_paragraph(meta["formule"])
    sg = doc.add_paragraph(style="Maison Signature"); sg.add_run(meta.get("signature") or d["identite"]["signature"])
    doc.add_paragraph(d["identite"]["raison_sociale"])
    if meta.get("annexes"):
        doc.add_paragraph(f"{L['annexes']} : " + ", ".join(meta["annexes"]), style="Maison Petit")
    W.remplacer_marqueurs(doc, marqueurs(d, meta))
    _proprietes(doc, d, meta)
    doc.save(str(dest))
    return []


def pv(meta, blocs, dest, d, sources=None):
    lg = meta.get("langue", "fr"); L = lib(lg)
    doc = ouvrir_gabarit("pv-assemblee.docx")
    W.vider_corps(doc)
    run = ajouter_runs(d)
    doc.add_paragraph(L["pv"], style="Title")
    p = doc.add_paragraph(style="Subtitle"); run(p, meta.get("titre", ""))
    W.tableau(doc, d, [L["societe"], meta.get("societe") or meta.get("client_nom", "")],
              [[L["datelieu"], f"{date_longue(meta.get('date_seance') or meta.get('date'), lg)} · {meta.get('lieu', '')}"],
               [L["pres"], meta.get("president", "")], [L["secr"], meta.get("secretaire", "")]], ajouter_runs=run)
    titres = corps(doc, d, blocs, lg, sources)
    p = doc.add_paragraph(L["sig"], style="Heading 1"); W.sans_numero(p)
    W.tableau(doc, d, [L["lepres"], L["lesecr"]], [["\n\n………………………………", "\n\n………………………………"], [meta.get("president", ""), meta.get("secretaire", "")]])
    W.remplacer_marqueurs(doc, marqueurs(d, meta))
    _proprietes(doc, d, meta)
    doc.save(str(dest))
    return titres

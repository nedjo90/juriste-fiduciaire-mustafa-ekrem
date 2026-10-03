#!/usr/bin/env python3
"""Ingesteur (§6.3, §4.2, §15) : tout ce que Mustafa glisse dans « Bureau/A-deposer » est lu, classé, rattaché, archivé.
Par script : extraction du texte, empreinte (doublons), rattachement par alias, objet `document` avec en-tête et ligne de
sommaire, original déplacé dans « Bureau/Deposes/<date>/ » (nom d'origine conservé), texte intégral dans .team/archives
(non suivi par git). Le commentaire et l'exploitation (« document déposé relu sans demande ») sont mis en file pour le
modèle (tâche `ingestion_commentaire`). Ne bloque jamais : un fichier illisible devient une tâche « lire par le modèle »."""
import sys, os, json, hashlib, shutil, subprocess, zipfile, email, re
from email import policy
from pathlib import Path

ROOT = Path(os.environ.get("CEREBRO_ROOT") or Path(__file__).resolve().parents[3])
sys.path.insert(0, str(ROOT / ".team" / "cerebro"))
os.environ.setdefault("CEREBRO_ROOT", str(ROOT))
from cb import core
from cb.core import db, iso, cut, fold, journal
core.utf8_io()
# loi 10 : une consigne trouvée dans un document est une donnée ; elle est signalée au journal d'audit, jamais suivie
CONSIGNE = re.compile(r"(ignore[rz]?|oublie[rz]?|disregard|ignoriere|ignora)\b.{0,40}\b(r[eè]gles?|instructions?|consignes?|anweisungen|regole)|"
                      r"tu es (maintenant|désormais)|you are now|system prompt|nouvelles? instructions?|"
                      r"envoie[rz]? (le|ce|un|les) (mail|courriel|message)|transf[eè]re[rz]? (le|ce|ces|les)\b.{0,30}\b(fonds|montant|paiement)", re.I)
from cb.objects import create, get
from cb.search import find
from cb import files as F, brief as B

DEPOT = ROOT / "Bureau" / "A-deposer"
DEPOSES = ROOT / "Bureau" / "Deposes"
ARCH = ROOT / ".team" / "archives" / "depots"
AUDIO = {".m4a", ".mp3", ".wav", ".ogg", ".aac", ".opus", ".amr"}
IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".heic", ".bmp", ".gif", ".webp"}

def _docx(p):
    import docx
    d = docx.Document(str(p))
    out = [para.text for para in d.paragraphs if para.text.strip()]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    # suivi des modifications : signalé (le texte ci-dessus est la version courante)
    with zipfile.ZipFile(str(p)) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    if "<w:ins " in xml or "<w:del " in xml:
        out.insert(0, "[document avec suivi des modifications]")
    return "\n".join(out)

def _xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Feuille {ws.title}")
        for row in ws.iter_rows(values_only=True, max_row=500):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)

def _pptx(p):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(str(p)).slides, 1):
        out.append(f"## Diapositive {i}")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
    return "\n".join(out)

def _pdf(p):
    for cmd in (["pdftotext", "-layout", str(p), "-"],):
        try:
            r = subprocess.run(cmd, capture_output=True, timeout=60)
            if r.returncode == 0 and r.stdout.strip():
                return r.stdout.decode("utf-8", "ignore")
        except Exception:
            pass
    try:
        import pypdf
        return "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(str(p)).pages)
    except Exception:
        return ""

def _eml(p):
    m = email.message_from_bytes(p.read_bytes(), policy=policy.default)
    body = m.get_body(preferencelist=("plain", "html"))
    txt = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        txt = re.sub(r"<[^>]+>", " ", txt)
    pj = [a.get_filename() for a in m.iter_attachments() if a.get_filename()]
    return f"De: {m['from']}\nÀ: {m['to']}\nDate: {m['date']}\nObjet: {m['subject']}\nPièces: {', '.join(pj) or '-'}\n\n{txt}"

CONNECTEURS = ROOT / ".team" / "scripts" / "connectors"
if str(CONNECTEURS) not in sys.path:
    sys.path.insert(0, str(CONNECTEURS))
PIECES_JOINTES = {}   # chemin du .msg → pièces jointes extraites (ingérées ensuite comme documents liés)
TRANSCRIPTIONS = {}   # chemin de l'audio → métadonnées de transcription (langue, modèle, partiel)

def _audio(p):
    """transcription locale (faster-whisper) ; sinon nature « audio_a_transcrire » (repli : question à Mustafa)"""
    try:
        import transcription as TR
        r = TR.transcrire(p)
        TRANSCRIPTIONS[str(p)] = r
        tete = f"[transcription automatique · langue {r['langue']} · modèle {r['modele']}{' · partielle' if r['partiel'] else ''}]\n\n"
        return tete + r["texte"], "audio"
    except Exception as e:
        TRANSCRIPTIONS[str(p)] = {"erreur": repr(e)[:300]}
        journal("ingester", fichier=p.name, transcription=repr(e)[:300])
        return "", "audio_a_transcrire"

def _msg(p):
    import outlook_msg
    h = hashlib.sha256(p.read_bytes()).hexdigest()[:16]
    texte, pieces = outlook_msg.lire_msg(p, ARCH / iso() / f"{h}-pieces")
    PIECES_JOINTES[str(p)] = pieces
    return texte

def _ancien_office(p):
    """ancien format Office : conversion par Microsoft Office du poste (COM, Windows), puis extracteurs modernes"""
    import office_read as office_lecture
    q = office_lecture.vers_moderne(p, ARCH / "_conv")
    if not q:
        return "", "a_lire_par_modele"
    try:
        t = {".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx}[q.suffix.lower()](q)
        return (t, "texte") if t.strip() else ("", "a_lire_par_modele")
    finally:
        try:
            q.unlink()
        except Exception:
            pass

def extraire(p):
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            return _docx(p), "texte"
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(p), "texte"
        if ext == ".pptx":
            return _pptx(p), "texte"
        if ext == ".pdf":
            t = _pdf(p)
            return (t, "texte") if len(t.strip()) > 40 else ("", "pdf_scanne")
        if ext == ".eml":
            return _eml(p), "mail"
        if ext in (".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm"):
            return core.lire(p), "texte"
        if ext == ".msg":
            return _msg(p), "mail"
        if ext in AUDIO:
            return _audio(p)
        if ext in IMAGES:
            return "", "image"
        if ext in (".doc", ".xls", ".ppt", ".rtf", ".odt", ".ods", ".odp", ".pps"):
            return _ancien_office(p)
    except Exception as e:
        journal("ingester", fichier=p.name, erreur=repr(e))
        if ext in AUDIO:  # le modèle ne lit pas l'audio : jamais de tâche « lecture par le modèle » pour une note vocale
            return "", "audio_a_transcrire"
    return "", "a_lire_par_modele"

def rattacher(texte, nom_fichier):
    """client probable par alias / recherche ; renvoie (client, confiance, autres objets cités)"""
    tete = fold(f"{nom_fichier} {texte[:300]}")
    corps = fold(texte[:20000])
    clients, cites = {}, []
    # seuls les alias et identifiants rattachent (un mot commun du texte ne suffit pas) ; alias long et en tête = plus fort
    for r in db().execute("SELECT a.alias_fold, a.id, o.type, o.client FROM alias a JOIN objets o ON o.id=a.id WHERE o.statut!='archive'"):
        a = r["alias_fold"]
        if len(a) < 4 or not re.search(r"\b" + re.escape(a) + r"\b", corps + " " + tete):
            continue
        cid = r["id"] if r["type"] == "client" else r["client"]
        if not cid:
            continue
        poids = len(a) * (2 if re.search(r"\b" + re.escape(a) + r"\b", tete) else 1)
        clients[cid] = clients.get(cid, 0) + poids
        if r["type"] != "client" and r["id"] not in cites:
            cites.append(r["id"])
    for i in core.ID_RE.findall(texte):
        o = get(i)
        if o and (o["client"] or o["type"] == "client"):
            cid = o["client"] or o["id"]
            clients[cid] = clients.get(cid, 0) + 30
    if not clients:
        return None, 0.0, []
    best = max(clients.items(), key=lambda x: x[1])
    total = sum(clients.values())
    return best[0], round(best[1] / total, 2), cites[:8]

def ingerer_fichier(p, parent=None):
    data = p.read_bytes()
    h = hashlib.sha256(data).hexdigest()[:16]
    con = db()
    dup = con.execute("SELECT id FROM objets WHERE type='document' AND data LIKE ?", (f'%"empreinte": "{h}"%',)).fetchone()
    jour = iso()
    dest = DEPOSES / jour
    dest.mkdir(parents=True, exist_ok=True)
    nom = p.name if len(p.name) <= 120 else p.stem[:110] + p.suffix  # chemins Windows longs (MAX_PATH)
    cible = dest / nom
    i = 2
    while cible.exists():
        cible = dest / f"{Path(nom).stem} ({i}){p.suffix}"; i += 1
    if dup:
        shutil.move(str(p), str(cible))
        journal("ingester", fichier=p.name, doublon_de=dup[0])
        return {"fichier": p.name, "doublon_de": dup[0]}
    texte, nature = extraire(p)
    consigne = CONSIGNE.search(texte or "")
    client, conf, cites = rattacher(texte, p.stem) if texte else (None, 0.0, [])
    parent_o = get(parent) if parent else None
    if parent_o and parent_o.get("client") and (not client or conf < 0.7):
        client, conf = parent_o["client"], 1.0  # pièce jointe : client du mail qui la porte
    a_confirmer = client is not None and conf < 0.7
    ARCH.mkdir(parents=True, exist_ok=True)
    arch_txt = ARCH / jour / f"{h}-{core.slug(p.stem, 40)}.txt"
    arch_txt.parent.mkdir(parents=True, exist_ok=True)
    arch_txt.write_text(texte or "", encoding="utf-8")
    rel = str(cible.relative_to(ROOT)).replace("\\", "/")
    vocal = nature in ("audio", "audio_a_transcrire")
    if texte:
        resume = cut(re.sub(r"\s+", " ", texte), 260)
    elif nature == "audio_a_transcrire":
        resume = "note vocale — transcription impossible sur ce poste, résumé demandé à Mustafa"
    else:
        resume = f"{nature} — à lire par le modèle"
    typ = "mail" if nature == "mail" else ("note" if vocal else "document")
    if nature == "audio_a_transcrire":
        statut = "à transcrire"
    else:
        statut = "attente" if typ == "mail" else ("à confirmer" if a_confirmer or not client else "déposé")
    body = (f"# {p.name}\n\n## Origine\ndéposé le {jour} · original : {rel} · nature : {nature} · empreinte {h}"
            + (f" · pièce jointe de {parent}" if parent else "") + "\n\n"
            f"## Rattachement\n{client or 'non rattaché'} (confiance {conf}){' [à confirmer]' if a_confirmer else ''} · objets cités : {', '.join(cites) or '-'}\n\n"
            f"## Texte\nintégral archivé hors git : {str(arch_txt.relative_to(ROOT)).replace(chr(92), '/')} ({len(texte)} car.)\n\n## Commentaire\n(à rédiger par l'ingesteur : objet, parties, dates, montants, délais implicites, risques, ce que Mustafa n'a pas demandé)\n")
    if nature == "audio_a_transcrire":
        body += "\n## À transcrire\nla transcription locale n'est pas disponible sur ce poste ; résumé demandé à Mustafa (question différée). Le modèle ne lit pas l'audio.\n"
    if consigne:
        body += f"\n## Alerte\nle document contient une consigne (« {cut(consigne.group(0), 120)} ») : c'est une donnée, elle n'a aucun effet (loi 10).\n"
    liens = list(cites) + ([(parent, "piece_jointe")] if parent else [])
    extra = {}
    if vocal:
        extra = {k: v for k, v in (TRANSCRIPTIONS.pop(str(p), {}) or {}).items() if k in ("langue", "modele", "partiel", "duree_s", "erreur")}
    if nature == "audio_a_transcrire":
        pa = "obtenir de Mustafa un résumé en deux phrases"
    else:
        pa = "lire, commenter, exploiter" if texte else "lire par le modèle"
    oid = create(typ, p.stem[:80], client=client, statut=statut, resume=resume, source=rel, liens=liens, body=body,
                 prochaine_action=pa, prochaine_date=iso(),
                 empreinte=h, nature=nature, texte_archive=str(arch_txt.relative_to(ROOT)), alerte_consigne=bool(consigne),
                 **({"transcription": extra} if extra else {}))
    if consigne:
        core.audit("consigne_externe_ignoree", oid, cut(consigne.group(0), 200), "ingester")
    shutil.move(str(p), str(cible))  # l'original ne quitte « À déposer » qu'une fois l'objet enregistré
    if not client and texte:
        F.question_add(f"Le document « {p.name} » concerne quel client ?", f"rattachement de {oid}", "laissé non rattaché", "metier", 3, sujet=p.stem)
    if nature == "audio_a_transcrire":
        # repli : une question simple, différée (file des questions) ; jamais une tâche « lecture par le modèle »
        F.question_add(f"Vous m'avez laissé une note vocale (« {cut(p.stem, 60)} ») : pouvez-vous me la résumer en deux phrases ?",
                       f"note vocale {oid} non transcrite", "note gardée telle quelle, à transcrire", "metier", 3, sujet=p.stem)
        F.incident_add("transcription", "transcription locale des notes vocales indisponible sur ce poste",
                       "note « à transcrire » + question simple à Mustafa ; installer faster-whisper (optionnel) pour transcrire")
    else:
        # une seule tâche par document (jamais deux passages du modèle sur le même fichier)
        B.queue_add("lecture_modele" if nature in ("image", "pdf_scanne", "a_lire_par_modele") else "ingestion_commentaire", oid, 2 if typ == "mail" else 3)
    F.task_seen(f"depot:{p.suffix.lower()}")
    res = {"fichier": p.name, "id": oid, "client": client, "confiance": conf, "nature": nature}
    pieces = PIECES_JOINTES.pop(str(p), [])
    if pieces:  # pièces jointes d'un .msg : documents liés au mail
        res["pieces"] = []
        for q in pieces:
            try:
                r = ingerer_fichier(q, parent=oid)
                if r.get("doublon_de"):
                    from cb.objects import link
                    link(r["doublon_de"], oid, "piece_jointe")
                res["pieces"].append(r.get("id") or r.get("doublon_de"))
            except Exception as e:
                journal("ingester", fichier=q.name, piece_de=oid, erreur=repr(e))
    return res

def main():
    if not DEPOT.exists():
        print(json.dumps({"deposes": 0}))
        return
    res = []
    for p in sorted(DEPOT.rglob("*")):
        if p.is_file() and not p.name.startswith((".", "~$")) and p.name not in ("desktop.ini", "Thumbs.db"):
            try:
                res.append(ingerer_fichier(p))
            except Exception as e:
                journal("ingester", fichier=p.name, erreur=repr(e))
                F.incident_add("ingestion", f"fichier illisible : {p.name}", "laissé dans A-deposer, nouvel essai au prochain cycle")
    for d in sorted([d for d in DEPOT.rglob("*") if d.is_dir()], reverse=True):
        try:
            d.rmdir()
        except OSError:
            pass
    print(json.dumps({"deposes": len(res), "resultats": res}, ensure_ascii=False))

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        journal("ingester", erreur=repr(e))
        print(json.dumps({"erreur": repr(e)}))

#!/usr/bin/env python3
"""Tests du producteur, du système de design, des gabarits et des portes (T-051, T-060 ; §17 critères 5, 19, 23, 31, 32, 34).
Racine jetable (copie de .team/cerebro, config, design) + dossier fictif : la base réelle n'est jamais touchée.
Usage : python .team/tests/test_production.py   → une ligne OK/ÉCHEC par vérification, puis le bilan."""
import os, sys, json, shutil, subprocess, tempfile, re
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
EQ = REPO / ".team"
PROD = EQ / "scripts" / "producer"
PORTES = EQ / "scripts" / "gates"
FIX = EQ / "tests" / "fixtures"
RES = []


def check(nom, cond, detail=""):
    RES.append(bool(cond))
    print(("OK     " if cond else "ÉCHEC  ") + nom + (f" — {detail}" if detail and not cond else ""))


def racine_jetable():
    t = Path(tempfile.mkdtemp(prefix="test-production-"))
    (t / ".team").mkdir()
    shutil.copytree(EQ / "cerebro", t / ".team" / "cerebro", ignore=shutil.ignore_patterns("*.db", "*.db-*", "__pycache__", "exports", "sauvegardes"))
    shutil.copytree(EQ / "config", t / ".team" / "config")
    (t / ".team" / "brain" / "firm" / "design").mkdir(parents=True)
    shutil.copy(EQ / "brain" / "firm" / "design" / "system.yaml", t / ".team" / "brain" / "firm" / "design")
    shutil.copy(EQ / "constitution.md", t / ".team" / "constitution.md")
    for d in ("Livrables", "Modeles", "Informatique"):
        (t / "Bureau" / d).mkdir(parents=True)
    return t


def env_de(t):
    e = dict(os.environ, CEREBRO_ROOT=str(t), CEREBRO_TODAY="2026-10-03", PRODUCTEUR_SANS_OUVERTURE="1", PRODUCTEUR_SANS_MMDC="1", PYTHONIOENCODING="utf-8")
    e.pop("CEREBRO_DB", None)
    return e


def py(t, *args, timeout=600):
    r = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True, encoding="utf-8", env=env_de(t), cwd=str(t), timeout=timeout)
    try:
        return json.loads(r.stdout.strip().splitlines()[-1]), r
    except Exception:
        return None, r


def main():
    t = racine_jetable()
    try:
        # --- données fictives + une source fictive (BIB-001) dans la racine jetable
        _, r = py(t, FIX / "fictional_case.py")
        check("dossier fictif chargé dans la racine jetable", r.returncode == 0, r.stderr[-300:])
        _, r = py(t, "-c", "import sys; sys.path.insert(0, '.team/cerebro'); from cb.objects import create; "
                           "print(create('source', '[FICTIF] texte de loi de test', resume='source fictive de test'))")
        check("source fictive BIB-001 créée", "BIB-001" in r.stdout, r.stderr[-300:])

        # --- système de design
        out, r = py(t, PROD / "design.py", "--verifier")
        check("contrastes de la palette conformes (WCAG)", out and out["ok"], r.stdout[-300:])

        # --- gabarits générés depuis le design
        out, r = py(t, PROD / "templates.py")
        noms = ["memo.docx", "lettre.docx", "pv-assemblee.docx", "modele-calcul.xlsx", "presentation.pptx", "gabarit-rapport.pdf"]
        check("six gabarits générés dans Bureau/Modeles", out and all((t / "Bureau" / "Modeles" / n).exists() for n in noms), r.stderr[-300:])
        from docx import Document
        g = Document(str(t / "Bureau" / "Modeles" / "memo.docx"))
        styles = {s.name for s in g.styles}
        check("gabarit mémo : styles de la maison et titres numérotés",
              {"Maison Encadré", "Maison Réserve"} <= styles and 'w:abstractNumId="90"' in g.part.numbering_part.element.xml)
        check("gabarit mémo : en-tête confidentialité et pied version/état du droit",
              "{{confidentialite}}" in g.sections[0].header.paragraphs[0].text and "{{pied}}" in g.sections[0].footer.paragraphs[0].text)

        # --- mémo de démonstration
        out, r = py(t, PROD / "produce.py", FIX / "memo_demo.md", "--role", "drafter")
        check("produce.py renvoie un JSON sans erreur", out and not out.get("erreur"), (r.stderr or r.stdout)[-500:])
        out_memo = out or {}
        docx_p = Path(out["principal"]) if out else None
        check("docx créé dans Bureau/Livrables/<client>/<date>-<objet>/",
              docx_p and docx_p.exists() and docx_p.parent.parent.parent == t / "Bureau" / "Livrables", str(docx_p))
        check("nom de fichier client-objet-date-vN (ASCII, kebab-case)",
              docx_p and re.fullmatch(r"[a-z0-9-]+-2026-10-03-v1\.docx", docx_p.name), docx_p and docx_p.name)
        doc = Document(str(docx_p))
        texte = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for tb in doc.tables for row in tb.rows for c in row.cells)
        check("docx issu du gabarit (styles maison, catégorie memo, marqueurs remplacés)",
              "Maison Réserve" in {s.name for s in doc.styles} and doc.core_properties.category == "memo" and "{{" not in texte
              and "Confidentiel" in doc.sections[0].header.paragraphs[0].text)
        check("pied de page : version et date d'état du droit", "État du droit au 1er" in doc.sections[0].footer.paragraphs[0].text,
              doc.sections[0].footer.paragraphs[0].text)
        check("PDF produit", out.get("pdf") and Path(out["pdf"]).exists() and Path(out["pdf"]).stat().st_size > 5000, str(out.get("pdf")))
        check("portes exécutées (9 portes, aucune non exécutée)", set(out["portes"]) >= {"liens", "sources", "typographie", "tics", "regle_zero", "presentation", "couverture", "visuel", "budget"},
              json.dumps(out["portes"]))
        check("citation de droit sans source → ⚠ inséré dans le livrable", re.search(r"35\s?%\s?⚠", texte), "")
        check("BIB-001 rendu en renvoi [1] + annexe des sources", "[1]" in texte and "[FICTIF] texte de loi de test" in texte)
        check("typographie corrigée : CHF 200000 → CHF 200'000.00", "CHF 200'000.00" in texte)
        check("graphique et schéma produits (PNG, SVG, draw.io)", any(f.endswith(".drawio") for f in out["fichiers"]) and any("graphique" in f and f.endswith(".png") for f in out["fichiers"]))
        check("livrable inscrit dans cerebro (LIV-…)", str(out.get("livrable", "")).startswith("LIV-"))
        check("ouverture automatique silencieuse en test", out.get("ouvert") is False)
        check("aucun fichier machine dans Bureau (source corrigée côté .team/run)", ".team" in out.get("source_corrigee", "") and not list((t / "Bureau").rglob("*.md")))
        out2, _ = py(t, PROD / "produce.py", FIX / "memo_demo.md", "--role", "drafter")
        check("deuxième production → version v2", out2 and out2["principal"].endswith("-v2.docx"))

        # --- portes en direct
        sys.path.insert(0, str(PORTES)); sys.path.insert(0, str(PROD))
        os.environ.update(env_de(t))
        import p_typo, p_tics, p_sources
        corr, _ = p_typo.corriger_texte("Montant : CHF 1234.5 le 2026-10-03", "fr")
        check("typographie : « CHF 1234.5 » → « CHF 1'234.50 »", "CHF 1'234.50" in corr.replace(" ", " "), corr)
        check("typographie FR : insécable avant « : » et date en toutes lettres", "Montant :" in corr and "3 octobre 2026" in corr, repr(corr))
        de, _ = p_typo.corriger_texte('Die Strasse „Weißhorn“ kostet CHF 12000.', "de")
        check("typographie DE : «…» suisses et ß → ss", "«Weisshorn»" in de and "CHF 12'000.00" in de, de)
        tics = p_tics.analyser("Merci pour votre message. N'hésitez pas à me contacter.", "fr", "mail")
        check("tic « n'hésitez pas » détecté", any(c["motif"] == "clôture passe-partout" for c in tics), json.dumps(tics, ensure_ascii=False))
        check("tic allemand « Zögern Sie nicht » détecté", p_tics.analyser("Zögern Sie nicht, uns anzurufen.", "de", "mail"))
        check("mention de l'IA et émoticône détectées", {"mention de l'IA", "émoticône"} <= {c["motif"] for c in p_tics.analyser("En tant qu'IA, je pense que oui 🙂", "fr")})
        uni = " ".join(["Le contrat prévoit une clause de durée fixe."] * 7)
        check("phrases de longueur uniforme détectées", any(c["motif"] == "phrases de longueur uniforme" for c in p_tics.analyser(uni, "fr")))
        m, n = p_sources.marquer_paragraphe("Le délai de réclamation est de 30 jours. La société a son siège à Échallens.")
        check("porte sources : ⚠ seulement sur la phrase de droit", n == 1 and "30 jours ⚠." in m and "Échallens." in m and m.count("⚠") == 1, m)
        m2, n2 = p_sources.marquer_paragraphe("Le taux est de 35 % (BIB-001).")
        check("porte sources : affirmation sourcée laissée intacte", n2 == 0 and "⚠" not in m2)

        # --- gates.py en ligne de commande sur un mail fautif (jamais bloquant : code 0)
        out3, r3 = py(t, PROD / "produce.py", FIX / "mail_demo.md", "--role", "drafter")
        check("mail : brouillon .eml et .txt produits", out3 and out3["principal"].endswith(".eml") and any(f.endswith(".txt") for f in out3["fichiers"]))
        check("mail : porte tics fermée, renvoyée à l'éditeur, livrable présenté avec réserves",
              out3["portes"].get("tics") == "ko" and out3["a_renvoyer"].get("tics") == "human-editor" and out3["reserves"])
        eml = Path(out3["principal"]).read_text(encoding="utf-8", errors="replace")
        check("mail : signature de la maison et marque de brouillon", "Mustafa Ekrem" in eml and "X-Unsent: 1" in eml)
        out4, r4 = py(t, PORTES / "gates.py", Path(out3["principal"]), "--role", "drafter")
        check("gates.py : JSON {porte: ok/ko, détails, corrections}, code 0", r4.returncode == 0 and out4 and all({"etat", "details", "corrections"} <= set(v) for v in out4["portes"].values()))

        # --- autres formats
        for nom, ext in (("lettre", ".docx"), ("pv", ".docx"), ("calcul", ".xlsx"), ("presentation", ".pptx")):
            o, rr = py(t, PROD / "produce.py", FIX / f"{nom}_demo.md")
            check(f"{nom} produit ({ext})", o and o.get("principal", "").endswith(ext), (rr.stderr or rr.stdout)[-300:])
            if nom == "calcul":
                check("calcul : aucune valeur en dur dans les formules (porte présentation)", o["portes"].get("presentation") == "ok")
            if nom == "presentation":
                check("présentation : titre-étiquette et notes manquantes signalés", o["portes"].get("presentation") == "ko")
        o, _ = py(t, PROD / "produce.py", FIX / "pv_demo.md")
        check("PV : structure complète reconnue", o and o["portes"].get("presentation") == "ok")

        # --- la valeur de firm.yaml remplace le défaut du design (raison sociale dans l'en-tête)
        _, rc = py(t, EQ / "cerebro" / "cerebro.py", "config", "set", "firm.raison_sociale", "Fiduciaire Exemple SA [FICTIF]", "--source", "test")
        o, _ = py(t, PROD / "produce.py", FIX / "lettre_demo.md")
        hd = Document(o["principal"]).sections[0].header.paragraphs[0].text if o else ""
        check("firm.yaml renseigné → raison sociale reprise dans l'en-tête", "Fiduciaire Exemple SA" in hd, hd)

        # --- règle zéro sur un texte pour Mustafa
        md = t / ".team" / "run" / "reponse.md"; md.parent.mkdir(parents=True, exist_ok=True)
        md.write_text("---\ntype: reponse\ndestinataire: mustafa\n---\nJ'ai relancé le script et mis à jour la base de données.\n", encoding="utf-8")
        o, _ = py(t, PORTES / "gates.py", md, "--portes", "regle_zero", "--sans-enregistrer")
        check("règle zéro : jargon détecté dans un texte pour Mustafa", o and o["portes"]["regle_zero"]["etat"] == "ko")

        # --- tableau de bord des principes
        o, _ = py(t, PORTES / "dashboard.py")
        check("tableau de bord : passages enregistrés par rôle", o and "drafter" in o["par_role"], json.dumps(o)[:300] if o else "")
        # Avec un moteur de rendu (Office, poppler, PyMuPDF) : images de contrôle. Sans : repli explicite, jamais bloquant.
        images = list((t / ".team" / "run" / "rendus").rglob("page-*.png"))
        repli = out_memo.get("portes", {}).get("visuel") in ("ok", "na") and any("contrôle visuel limité" in x for x in out_memo.get("reserves", []))
        check("visuel : images de contrôle rendues, ou repli « contrôle visuel limité » non bloquant", images or repli,
              json.dumps({"visuel": out_memo.get("portes", {}).get("visuel"), "reserves": out_memo.get("reserves", [])[:4]}, ensure_ascii=False))
    finally:
        if not os.environ.get("GARDER_RACINE"):
            shutil.rmtree(t, ignore_errors=True)
        else:
            print("racine gardée :", t)
    ok = all(RES)
    print(f"{'OK' if ok else 'ÉCHEC'} — {sum(RES)}/{len(RES)} vérifications")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())

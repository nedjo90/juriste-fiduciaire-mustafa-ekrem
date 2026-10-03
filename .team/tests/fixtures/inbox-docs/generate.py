#!/usr/bin/env python3
"""Données de répétition : vrais fichiers fictifs pour « À déposer » (§3, critères 18 et 23). Tout est inventé
(dossier fictif Rochat / Lémantech). Régénération idempotente : python generate.py [--dest DOSSIER]
Fichiers : mail .eml (avec pièce jointe), mail Outlook .msg (fichier composé OLE écrit ici, sans Outlook, relu par
extract-msg), PDF texte, PDF image (scan simulé, sans couche texte), .docx avec suivi des modifications réel (w:ins /
w:del), .xlsx avec formules, note vocale .wav (voix de synthèse hors ligne espeak-ng si présent ; sinon absente, notée)."""
import sys, os, io, struct, shutil, subprocess, datetime as dt, argparse, math
from pathlib import Path
from email.message import EmailMessage
from email.utils import format_datetime

ICI = Path(__file__).resolve().parent
FICTIF = "[FICTIF]"


# ------------------------------------------------------------------ .eml
def eml(dest, pdf_joint=None):
    m = EmailMessage()
    m["From"] = "Jean-Marc Rochat <jm.rochat@exemple.ch>"
    m["To"] = "Mustafa Ekrem <m.ekrem@fiduciaire-exemple.ch>"
    m["Subject"] = "Décision de taxation 2025 de Rochat Holding SA"
    m["Date"] = format_datetime(dt.datetime(2026, 9, 24, 9, 12, tzinfo=dt.timezone(dt.timedelta(hours=2))))
    m["Message-ID"] = "<fictif-0001@exemple.ch>"
    m.set_content(f"""{FICTIF} Bonjour Monsieur Ekrem,

Nous avons reçu le 22 septembre la décision de taxation 2025 de Rochat Holding SA (impôt cantonal et communal et impôt
fédéral direct), envoyée par l'Administration cantonale des impôts du canton de Vaud. Le bénéfice imposable retenu est
de CHF 48'250 de plus que notre déclaration : la provision pour le litige avec le fournisseur a été refusée.

Pouvons-nous contester ? Jusqu'à quand ? Je vous joins la décision.

Meilleures salutations
Jean-Marc Rochat
Rochat Holding SA, Échallens
""")
    if pdf_joint and Path(pdf_joint).exists():
        m.add_attachment(Path(pdf_joint).read_bytes(), maintype="application", subtype="pdf", filename=Path(pdf_joint).name)
    (dest / "mail-rochat-taxation-2025.eml").write_bytes(bytes(m))


# ------------------------------------------------------------------ .msg (fichier composé OLE / CFB v3)
FREE, ENDCHAIN, FATSECT, NOSTREAM = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD, 0xFFFFFFFF


def _cfb(entrees):
    """entrees : {"chemin/vers/flux": bytes} ; stockages déduits des chemins. Renvoie les octets du fichier composé.
    Flux < 4096 octets dans le mini-flux (secteurs de 64), les autres en secteurs de 512 (spécification MS-CFB)."""
    SS, MSS, CUT = 512, 64, 4096
    arbre = {"": {"type": 5, "enfants": {}, "data": b""}}
    for chemin, data in entrees.items():
        parts = chemin.split("/")
        cur = ""
        for i, p in enumerate(parts):
            nxt = f"{cur}/{p}" if cur else p
            if nxt not in arbre:
                arbre[nxt] = {"type": 2 if i == len(parts) - 1 else 1, "enfants": {}, "data": data if i == len(parts) - 1 else b""}
                arbre[cur]["enfants"][p] = nxt
            cur = nxt
    ordre = [""] + sorted(k for k in arbre if k)
    idx = {k: i for i, k in enumerate(ordre)}
    # mini-flux
    mini, minifat = bytearray(), []
    big = []
    for k in ordre:
        e = arbre[k]
        if e["type"] != 2:
            continue
        d = e["data"]
        if len(d) < CUT:
            n = max(1, math.ceil(len(d) / MSS)) if d else 0
            e["start"] = len(minifat) if n else ENDCHAIN
            for j in range(n):
                minifat.append(len(minifat) + 1 if j < n - 1 else ENDCHAIN)
            mini += d + b"\0" * (n * MSS - len(d))
        else:
            big.append(k)
    secteurs, fat = [], []

    def ajouter(data):
        n = math.ceil(len(data) / SS)
        debut = len(secteurs)
        for j in range(n):
            secteurs.append(data[j * SS:(j + 1) * SS].ljust(SS, b"\0"))
            fat.append(debut + j + 1 if j < n - 1 else ENDCHAIN)
        return debut if n else ENDCHAIN

    for k in big:
        arbre[k]["start"] = ajouter(arbre[k]["data"])
    root_start = ajouter(bytes(mini)) if mini else ENDCHAIN
    mf = b"".join(struct.pack("<I", x) for x in minifat)
    minifat_start = ajouter(mf) if mf else ENDCHAIN
    n_minifat = math.ceil(len(mf) / SS)

    # répertoire : arbres binaires équilibrés par stockage (comparaison : longueur puis majuscules)
    gauche, droite, enfant = {}, {}, {}

    def cle(nom):
        return (len(nom), nom.upper())

    def equilibrer(noms):
        if not noms:
            return NOSTREAM
        m = len(noms) // 2
        racine = noms[m]
        gauche[racine] = equilibrer(noms[:m])
        droite[racine] = equilibrer(noms[m + 1:])
        return idx[racine]

    for k in ordre:
        enf = sorted(arbre[k]["enfants"].values(), key=lambda c: cle(c.split("/")[-1]))
        enfant[k] = equilibrer(enf)
    rep = bytearray()
    for k in ordre:
        e = arbre[k]
        nom = "Root Entry" if k == "" else k.split("/")[-1]
        n16 = nom.encode("utf-16-le") + b"\0\0"
        start = root_start if k == "" else e.get("start", 0) if e["type"] == 2 else 0
        taille = len(mini) if k == "" else len(e["data"]) if e["type"] == 2 else 0
        rep += struct.pack("<64sHBBIII16sIQQIQ", n16.ljust(64, b"\0"), len(n16), e["type"], 1,
                           gauche.get(k, NOSTREAM), droite.get(k, NOSTREAM), enfant.get(k, NOSTREAM),
                           b"\0" * 16, 0, 0, 0, start if start != ENDCHAIN or k == "" else ENDCHAIN, taille)
    while len(rep) % SS:
        rep += struct.pack("<64sHBBIII16sIQQIQ", b"\0" * 64, 0, 0, 0, NOSTREAM, NOSTREAM, NOSTREAM, b"\0" * 16, 0, 0, 0, 0, 0)
    rep_start = ajouter(bytes(rep))
    # FAT (elle se compte elle-même)
    n_fat = 1
    while (len(secteurs) + n_fat) > n_fat * (SS // 4):
        n_fat += 1
    fat_debut = len(secteurs)
    fat += [FATSECT] * n_fat
    fat += [FREE] * (n_fat * (SS // 4) - len(fat))
    fatb = b"".join(struct.pack("<I", x) for x in fat)
    for j in range(n_fat):
        secteurs.append(fatb[j * SS:(j + 1) * SS])
    assert n_fat <= 109
    difat = [fat_debut + j for j in range(n_fat)] + [FREE] * (109 - n_fat)
    tete = struct.pack("<8s16sHHHHH6sIIIIIIIII", bytes.fromhex("D0CF11E0A1B11AE1"), b"\0" * 16, 0x003E, 0x0003, 0xFFFE, 9, 6,
                       b"\0" * 6, 0, n_fat, rep_start, 0, CUT, minifat_start, n_minifat, ENDCHAIN, 0)
    tete += b"".join(struct.pack("<I", x) for x in difat)
    return tete + b"".join(secteurs)


def _u(s):
    return s.encode("utf-16-le")


def _filetime(d):
    return int((d - dt.datetime(1601, 1, 1, tzinfo=dt.timezone.utc)).total_seconds() * 10_000_000)


def msg(dest):
    """mail Outlook (.msg) conforme à MS-OXMSG : propriétés en flux __substg1.0_XXXXTTTT, flux des propriétés fixes,
    un destinataire, une pièce jointe ; chaînes en UTF-16 (type 001F)."""
    quand = dt.datetime(2026, 9, 29, 16, 40, tzinfo=dt.timezone.utc)
    corps = (f"{FICTIF} Bonjour Monsieur Ekrem,\r\n\r\nPour la convention d'actionnaires de Lémantech Sàrl, les deux investisseurs "
             "souhaitent un droit de préemption, une clause d'entraînement (drag-along) et un siège au conseil des gérants. "
             "Signature visée fin octobre 2026. Pouvez-vous nous envoyer un premier projet d'ici le 9 octobre ?\r\n\r\n"
             "Avec mes meilleures salutations\r\nLuca Rochat\r\nGérant, Lémantech Sàrl, Carouge\r\n")
    chaines = {0x001A: "IPM.Note", 0x0037: "Convention d'actionnaires Lémantech — premier projet", 0x1000: corps,
               0x0C1A: "Luca Rochat", 0x0C1F: "luca.rochat@lemantech-exemple.ch", 0x0C1E: "SMTP",
               0x0042: "Luca Rochat", 0x0065: "luca.rochat@lemantech-exemple.ch", 0x0E04: "Mustafa Ekrem", 0x0E1D: "Convention d'actionnaires Lémantech — premier projet",
               0x0070: "Convention d'actionnaires Lémantech — premier projet"}
    E = {}
    props = bytearray(b"\0" * 32)  # en-tête du message (8 réservés, next recip id, next attach id, recip count, attach count, 8 réservés)
    struct.pack_into("<IIII", props, 8, 1, 1, 1, 1)
    for pid, s in chaines.items():
        d = _u(s)
        E[f"__substg1.0_{pid:04X}001F"] = d
        props += struct.pack("<IIII", (pid << 16) | 0x001F, 6, len(d) + 2, 0)
    for pid in (0x0E06, 0x0039):  # MessageDeliveryTime, ClientSubmitTime (FILETIME)
        props += struct.pack("<IIQ", (pid << 16) | 0x0040, 6, _filetime(quand))
    props += struct.pack("<IIQ", (0x0E07 << 16) | 0x0003, 6, 1)  # MessageFlags : lu
    E["__properties_version1.0"] = bytes(props)
    # destinataire
    r = "__recip_version1.0_#00000000/"
    rp = bytearray(b"\0" * 8)
    for pid, s in {0x3001: "Mustafa Ekrem", 0x3003: "m.ekrem@fiduciaire-exemple.ch", 0x3002: "SMTP", 0x39FE: "m.ekrem@fiduciaire-exemple.ch"}.items():
        E[r + f"__substg1.0_{pid:04X}001F"] = _u(s)
        rp += struct.pack("<IIII", (pid << 16) | 0x001F, 6, len(_u(s)) + 2, 0)
    rp += struct.pack("<IIQ", (0x0C15 << 16) | 0x0003, 6, 1)  # RecipientType : To
    E[r + "__properties_version1.0"] = bytes(rp)
    # pièce jointe (texte)
    a = "__attach_version1.0_#00000000/"
    pj = (f"{FICTIF} Liste des points demandés par les investisseurs :\r\n1. droit de préemption\r\n2. drag-along / tag-along\r\n"
          "3. siège au conseil des gérants\r\n4. bad leaver / good leaver\r\n").encode("utf-8")
    ap = bytearray(b"\0" * 8)
    for pid, s in {0x3707: "points-investisseurs.txt", 0x3704: "points~1.txt", 0x3703: ".txt", 0x370E: "text/plain", 0x3001: "points-investisseurs.txt"}.items():
        E[a + f"__substg1.0_{pid:04X}001F"] = _u(s)
        ap += struct.pack("<IIII", (pid << 16) | 0x001F, 6, len(_u(s)) + 2, 0)
    E[a + "__substg1.0_37010102"] = pj
    ap += struct.pack("<IIII", (0x3701 << 16) | 0x0102, 6, len(pj), 0)
    ap += struct.pack("<IIQ", (0x3705 << 16) | 0x0003, 6, 1)  # AttachMethod : par valeur
    E[a + "__properties_version1.0"] = bytes(ap)
    # table des propriétés nommées (vide mais présente, exigée par MS-OXMSG)
    n = "__nameid_version1.0/"
    E[n + "__substg1.0_00020102"] = b""
    E[n + "__substg1.0_00030102"] = b""
    E[n + "__substg1.0_00040102"] = b""
    (dest / "mail-lemantech-convention.msg").write_bytes(_cfb(E))


# ------------------------------------------------------------------ PDF texte / PDF image
TEXTE_DECISION = [
    "Administration cantonale des impôts (FICTIF)",
    "Décision de taxation — période fiscale 2025",
    "Contribuable : Rochat Holding SA, Échallens (FICTIF)",
    "Impôt cantonal et communal et impôt fédéral direct",
    "Bénéfice déclaré : CHF 312'400 · Bénéfice imposable retenu : CHF 360'650",
    "Reprise : provision pour litige fournisseur CHF 48'250 (justificatifs insuffisants)",
    "Notifiée le 22 septembre 2026.",
    "Voies de droit : réclamation écrite et motivée auprès de l'autorité de taxation.",
]


def pdf_texte(dest):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    p = dest / "decision-taxation-2025-rochat-holding.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    c.setTitle("Décision de taxation 2025 (FICTIF)")
    y = 780
    for i, l in enumerate(TEXTE_DECISION):
        c.setFont("Helvetica-Bold" if i < 2 else "Helvetica", 13 if i < 2 else 11)
        c.drawString(60, y, l)
        y -= 28
    c.showPage(); c.save()
    return p


def pdf_image(dest):
    """scan simulé : texte rastérisé (Pillow), léger bruit et rotation, PDF sans couche texte"""
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    import random
    random.seed(7)
    W, H = 1654, 2339  # A4 à 200 dpi
    im = Image.new("L", (W, H), 250)
    d = ImageDraw.Draw(im)
    police = None
    for f in ("DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "arial.ttf", "C:/Windows/Fonts/arial.ttf", "/Library/Fonts/Arial.ttf"):
        try:
            police = ImageFont.truetype(f, 38)
            break
        except Exception:
            continue
    police = police or ImageFont.load_default()
    lignes = ["PROCÈS-VERBAL (FICTIF)", "Assemblée générale ordinaire de Rochat Holding SA", "Échallens, le 30 septembre 2026",
              "Capital représenté : 100 %", "1. Comptes annuels 2025 approuvés", "2. Dividende brut CHF 200'000, payable le 30.09.2026",
              "3. Décharge au conseil d'administration", "4. Renonciation au contrôle restreint confirmée", "Le président : J.-M. Rochat"]
    y = 220
    for l in lignes:
        d.text((180, y), l, fill=25, font=police)
        y += 110
    for _ in range(4000):
        d.point((random.randrange(W), random.randrange(H)), fill=random.randrange(120, 220))
    im = im.rotate(0.8, fillcolor=250, resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.6))
    p = dest / "pv-ag-2026-rochat-holding-scan.pdf"
    im.convert("RGB").save(str(p), "PDF", resolution=200.0)
    return p


# ------------------------------------------------------------------ .docx avec suivi des modifications
def docx_suivi(dest):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc = Document()
    doc.add_heading("Convention d'actionnaires de Lémantech Sàrl — projet (FICTIF)", 1)
    p = doc.add_paragraph("Art. 4 Droit de préemption. Chaque associé bénéficie d'un droit de préemption sur les parts cédées, ")
    quand = "2026-10-01T10:15:00Z"

    def run(texte, supprime=False):
        r = OxmlElement("w:r")
        t = OxmlElement("w:delText" if supprime else "w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = texte
        r.append(t)
        return r

    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "1"); ins.set(qn("w:author"), "Investisseur A (FICTIF)"); ins.set(qn("w:date"), quand)
    ins.append(run("à exercer dans un délai de trente jours dès la notification de la cession projetée, "))
    p._p.append(ins)
    dl = OxmlElement("w:del")
    dl.set(qn("w:id"), "2"); dl.set(qn("w:author"), "Investisseur A (FICTIF)"); dl.set(qn("w:date"), quand)
    dl.append(run("au prix fixé par le conseil des gérants.", supprime=True))
    p._p.append(dl)
    ins2 = OxmlElement("w:ins")
    ins2.set(qn("w:id"), "3"); ins2.set(qn("w:author"), "Investisseur A (FICTIF)"); ins2.set(qn("w:date"), quand)
    ins2.append(run("au prix offert par le tiers acquéreur de bonne foi."))
    p._p.append(ins2)
    doc.add_paragraph("Art. 5 Clause d'entraînement. Si des associés détenant au moins 75 % des parts acceptent une offre, ils peuvent "
                      "obliger les autres associés à céder leurs parts aux mêmes conditions.")
    doc.core_properties.author = "Mustafa Ekrem"
    doc.core_properties.title = "Convention d'actionnaires Lémantech (projet FICTIF)"
    p = dest / "convention-lemantech-suivi-modifications.docx"
    doc.save(str(p))
    return p


# ------------------------------------------------------------------ .xlsx
def xlsx(dest):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Hypothèses"
    ws.append(["Décompte TVA T3 2026 — Lémantech Sàrl (FICTIF)"])
    ws.append(["Chiffre d'affaires imposable (CHF)", 182400])
    ws.append(["Impôt préalable (CHF)", 6120.35])
    ws.append(["Taux normal (à lire dans la LTVA, source : bibliothèque)", None])
    c = wb.create_sheet("Calcul")
    c.append(["Poste", "Montant"])
    c.append(["TVA due", "=Hypothèses!B2*Hypothèses!B4"])
    c.append(["Impôt préalable", "=Hypothèses!B3"])
    c.append(["Solde à payer", "=B2-B3"])
    p = dest / "decompte-tva-t3-2026-lemantech.xlsx"
    wb.save(str(p))
    return p


# ------------------------------------------------------------------ .wav
def wav(dest):
    exe = shutil.which("espeak-ng") or shutil.which("espeak")
    p = dest / "note-vocale-rochat-succession.wav"
    texte = ("Note pour le dossier Rochat. Jean-Marc Rochat m'a appelé ce matin. Il veut transmettre la menuiserie à son fils "
             "Luca d'ici deux mille vingt-huit et demande si une donation des actions de la holding est possible. "
             "Rappeler avant le rendez-vous de jeudi.")
    if not exe:
        return None
    subprocess.run([exe, "-v", "fr", "-s", "150", "-w", str(p), texte], capture_output=True, timeout=120)
    return p if p.exists() and p.stat().st_size > 1000 else None


def generer(dest=ICI):
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    pdf = pdf_texte(dest)
    eml(dest, pdf)
    msg(dest)
    pdf_image(dest)
    docx_suivi(dest)
    xlsx(dest)
    w = wav(dest)
    return {"fichiers": sorted(f.name for f in dest.iterdir() if f.is_file() and f.suffix != ".py"), "wav": bool(w)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--dest", default=str(ICI))
    import json
    print(json.dumps(generer(ap.parse_args().dest), ensure_ascii=False))
